"""🔴 MEGA TEST 4 — Deliverable completeness (Phase 105).

Covers Parts VIII-IX (Phases 77-103). Nothing in Part X starts until this file
is green.

MEGA TEST 3 asked whether the results still say the same thing. This one asks
whether the deliverables built from them are all there, all traceable, and all
reproducible from the commands they record:

* T105.1 -- 30 tables, non-empty, each naming the file it was built from;
* T105.2 -- 35 graphs as 300 dpi PNG plus a CSV that is neither empty nor all-NaN;
* T105.3 -- 20 diagrams (SVG + PNG), 20 algorithms (DOCX + text), the equations reference;
* T105.4 -- every evidence row resolves, and five recorded commands, re-run, reproduce
  their artifact byte for byte (content digest, newline-normalized);
* T105.5 -- every number in the Q1 pack matches its source CSV;
* T105.6 -- everything not produced is in ``missing_outputs_report.txt`` with a reason.

T105.7 is [TEST/MANUAL] -- a human opens five tables and five graphs -- and is
not asserted here.
"""

from __future__ import annotations

import json
import os
import re
import shlex
import subprocess
import sys
from pathlib import Path

import pytest

from src.reporting import evidence_pack as ep
from src.reporting import q1_pack as q1
from src.utils.evidence import PROJECT_ROOT, read_evidence

pytestmark = pytest.mark.filterwarnings("ignore::FutureWarning")

OUTPUTS = PROJECT_ROOT / "outputs"
FIGURES = OUTPUTS / "13_figures_diagrams"
ALGORITHMS = OUTPUTS / "14_algorithms"
EXCLUDED = ("_checkpoints", "THESIS_ASSETS", "Q1_PAPER_ASSETS")

TABLE_IDS = ["T" + str(n).zfill(2) for n in range(1, 31)]
GRAPH_IDS = ["G" + str(n).zfill(2) for n in range(1, 36)]
DIAGRAM_IDS = ["F" + str(n).zfill(2) for n in range(1, 21)]
ALGORITHM_IDS = ["ALG-" + str(n).zfill(2) for n in range(1, 21)]

#: T105.4's five spot-checks: fast, deterministic commands spanning five phases.
SPOT_CHECK_IDS = ("LIT-01", "EQUATIONS-CSV", "T29", "ALG-20", "STAT-01")


def _gitignored(path: str) -> bool:
    result = subprocess.run(
        ["git", "check-ignore", "-q", path], cwd=PROJECT_ROOT, capture_output=True, check=False
    )
    return result.returncode == 0


def _one(pattern: str, folder: Path = OUTPUTS) -> Path:
    found = [p for p in folder.rglob(pattern) if not any(part in p.parts for part in EXCLUDED)]
    assert len(found) == 1, pattern + ": expected one file, found " + str(found)
    return found[0]


@pytest.fixture(scope="module")
def outputs_present() -> None:
    if not (FIGURES / "figure_registry.csv").is_file():
        pytest.skip("no generated deliverables to audit")


# ---------------------------------------------------------------------------
# T105.1 -- the thirty tables
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("table_id", TABLE_IDS)
def test_t105_1_table_exists_is_non_empty_and_names_its_source(
    outputs_present: None, table_id: str
) -> None:
    import pandas as pd

    meta_path = _one(table_id + "_*.meta.json")
    stem = meta_path.name[: -len(".meta.json")]
    csv_path = meta_path.parent / (stem + ".csv")
    docx_path = meta_path.parent / (stem + ".docx")
    assert csv_path.is_file() and docx_path.is_file(), table_id + ": CSV or DOCX missing"
    frame = pd.read_csv(csv_path)
    assert len(frame) > 0 and len(frame.columns) > 0, table_id + ": empty table"

    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    sources = [source["path"] for source in meta.get("sources", [])]
    assert sources, table_id + ": the table names no source file"
    for source in sources:
        exists = (PROJECT_ROOT / source).exists()
        assert exists or _gitignored(source), table_id + ": source not on disk: " + source

    from docx import Document

    document = Document(str(docx_path))
    assert document.tables and len(document.tables[0].rows) == len(frame) + 1, (
        table_id + ": the DOCX does not hold the CSV's rows"
    )


# ---------------------------------------------------------------------------
# T105.2 -- the thirty-five graphs
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("graph_id", GRAPH_IDS)
def test_t105_2_graph_is_a_300_dpi_png_with_a_non_empty_csv(
    outputs_present: None, graph_id: str
) -> None:
    import pandas as pd
    from PIL import Image

    from src.reporting.graphs import read_registry

    rows = {row["figure_id"]: row for row in read_registry(FIGURES / "figure_registry.csv")}
    assert graph_id in rows, graph_id + " is not registered"
    png = FIGURES / rows[graph_id]["filename"]
    data = FIGURES / rows[graph_id]["source_csv"]
    assert png.is_file() and data.is_file(), graph_id + ": PNG or CSV missing"

    dpi = Image.open(png).info.get("dpi")
    assert dpi is not None and all(abs(float(d) - 300) < 1 for d in dpi), (
        graph_id + ": dpi " + str(dpi)
    )

    frame = pd.read_csv(data)
    assert len(frame) > 0, graph_id + ": the CSV is empty"
    assert frame.notna().any().any(), graph_id + ": the CSV is all-NaN"
    numeric = frame.select_dtypes("number")
    assert not numeric.empty and numeric.notna().any().any(), graph_id + ": no plotted number"


# ---------------------------------------------------------------------------
# T105.3 -- diagrams, algorithms, equations
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("figure_id", DIAGRAM_IDS)
def test_t105_3_diagram_has_svg_and_png(outputs_present: None, figure_id: str) -> None:
    svg = _one(figure_id + "_*.svg", FIGURES)
    png = svg.with_suffix(".png")
    assert png.is_file(), figure_id + ": PNG missing"
    assert svg.stat().st_size > 0 and png.stat().st_size > 0


@pytest.mark.parametrize("alg_id", ALGORITHM_IDS)
def test_t105_3_algorithm_has_docx_and_text(outputs_present: None, alg_id: str) -> None:
    txt = _one(alg_id + "_*.txt", ALGORITHMS)
    assert txt.with_suffix(".docx").is_file(), alg_id + ": DOCX missing"
    assert txt.read_text(encoding="utf-8").strip(), alg_id + ": empty text"


def test_t105_3_equations_reference_is_present(outputs_present: None) -> None:
    for suffix in (".docx", ".tex", ".csv"):
        assert (ALGORITHMS / ("equations_reference" + suffix)).is_file(), suffix


# ---------------------------------------------------------------------------
# T105.4 -- the evidence index, and five commands re-run
# ---------------------------------------------------------------------------


def test_t105_4_every_evidence_row_resolves_to_a_real_file(outputs_present: None) -> None:
    missing = [
        row["evidence_id"] + " -> " + row["filename"]
        for row in read_evidence()
        if ep.row_status(row)[0] == "failed" and not _gitignored(row["filename"])
    ]
    assert not missing, missing


@pytest.mark.slow
@pytest.mark.parametrize("evidence_id", SPOT_CHECK_IDS)
def test_t105_4_recorded_command_reproduces_its_artifact(
    outputs_present: None, evidence_id: str, tmp_path: Path
) -> None:
    from src.reporting.tables import content_digest

    row = next((r for r in read_evidence() if r["evidence_id"] == evidence_id), None)
    assert row is not None and row["command"], evidence_id + " has no recorded command"
    for source in ep.source_paths(row["source_data"]):
        if not (PROJECT_ROOT / source).exists():
            pytest.skip(evidence_id + ": input not present (" + source + ")")

    artifact = PROJECT_ROOT / row["filename"]
    before = content_digest(artifact)[0]

    argv = shlex.split(row["command"])
    assert argv[0] == "python", "commands are recorded as python invocations"
    env = dict(os.environ)
    # A re-run records itself; point that at scratch, never the committed manifest.
    env["HEARTGUARD__PATHS__OUTPUTS__RUN_MANIFEST"] = str(tmp_path / "run_manifest.json")
    result = subprocess.run(
        [sys.executable, *argv[1:]],
        cwd=PROJECT_ROOT,
        env=env,
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, evidence_id + ": command failed\n" + result.stderr[-2000:]
    assert artifact.is_file(), evidence_id + ": the command did not rewrite its artifact"
    assert content_digest(artifact)[0] == before, (
        evidence_id + ": re-running `" + row["command"] + "` changed " + row["filename"]
    )


# ---------------------------------------------------------------------------
# T105.5 -- the Q1 pack against its sources
# ---------------------------------------------------------------------------


def test_t105_5_every_q1_number_matches_its_source(outputs_present: None) -> None:
    if not (q1.q1_dir() / "Q1_results_narrative.docx").is_file():
        pytest.skip("Q1 pack not generated")
    problems = {asset: issues for asset, issues in q1.verify_q1_pack().items() if issues}
    assert not problems, problems


# ---------------------------------------------------------------------------
# T105.6 -- the missing-outputs report
# ---------------------------------------------------------------------------


def test_t105_6_every_unproduced_mandatory_item_is_reported_with_a_reason(
    outputs_present: None,
) -> None:
    results = ep.check_mandatory()
    failed = [
        item["item_id"] + ": " + item["detail"] for item in results if item["status"] == "failed"
    ]
    assert not failed, "mandatory items missing with no reason:\n" + "\n".join(failed)

    report = ep.missing_report_path().read_text(encoding="utf-8")
    assert ep.REPORT_BEGIN in report, "the generated completeness block is missing"
    block = report.split(ep.REPORT_BEGIN, 1)[1].split(ep.REPORT_END, 1)[0]
    for item in results:
        if item["status"] == "not produced":
            assert "] " + item["item_id"] + " -- " in block, item["item_id"]
            assert "Reason: " + item["detail"] in block, item["item_id"] + " has no reason"


def test_t105_6_every_skipped_task_has_a_missing_outputs_entry() -> None:
    todo = PROJECT_ROOT / "Docs" / "todo.md"
    if not todo.is_file():
        pytest.skip("Docs/todo.md is not in this checkout")
    skipped = re.findall(r"^- \[-\] \*\*(T\d+\.\d+)\*\*", todo.read_text(encoding="utf-8"), re.M)
    report = ep.missing_report_path().read_text(encoding="utf-8")
    unreported = [task for task in skipped if task not in report]
    assert not unreported, "skipped tasks with no missing-outputs entry: " + ", ".join(unreported)
