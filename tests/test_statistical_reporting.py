"""T83.7 -- the statistical reporting gate.

The gate: **the significance matrix and the summary DOCX exist, the
multiple-comparison correction is applied, and any non-significant ensemble
advantage is stated plainly.**

The last clause is the one that needs a real assertion rather than a file-exists
check. It is satisfied only if the findings document actually names the
comparisons where M6/M7 did not beat a simpler model -- and separately, and
without conflating the two, the comparisons where the test could not have
detected a difference at all.
"""

from __future__ import annotations

from typing import Any

import numpy as np
import pytest

SECTION = "outputs/12_statistics"


def _root() -> Any:
    from src.utils.evidence import PROJECT_ROOT

    return PROJECT_ROOT / SECTION


def _csv(name: str) -> Any:
    import pandas as pd

    path = _root() / name
    if not path.is_file():
        pytest.skip(str(path) + " does not exist; run scripts/35_statistical_reporting.py")
    return pd.read_csv(path)


def _findings() -> str:
    path = _root() / "statistical_findings.md"
    if not path.is_file():
        pytest.skip(str(path) + " does not exist; run scripts/35_statistical_reporting.py")
    return path.read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# clause 1 -- the matrix and the DOCX exist
# ---------------------------------------------------------------------------


def test_the_significance_matrix_exists_and_is_a_real_matrix() -> None:
    matrix = _csv("statistical_significance_matrix.csv")
    assert len(matrix)
    assert "model" in matrix.columns
    assert {"n", "n_description", "p_value_kind"} <= set(matrix.columns)
    assert set(matrix["p_value_kind"]) == {"p_holm"}

    models = sorted(set(matrix["model"]))
    for model in models:
        assert model in matrix.columns, model + " has a row but no column"

    # Symmetric, with an empty diagonal: a model was not compared with itself,
    # and a 1.0 there would read as a test that found no difference.
    block = matrix[(matrix["run"] == "EXP-A2") & (matrix["metric"] == "sensitivity")]
    if len(block):
        for _, row in block.iterrows():
            assert np.isnan(float(row[str(row["model"])])), "the diagonal must be empty"
        for _, row in block.iterrows():
            for other in models:
                if other == row["model"] or other not in block.columns:
                    continue
                mirror = block[block["model"] == other]
                if not len(mirror):
                    continue
                a = float(row[other])
                b = float(mirror.iloc[0][str(row["model"])])
                assert np.isclose(a, b, equal_nan=True), (row["model"], other)


def test_the_summary_docx_exists_and_holds_the_four_test_families() -> None:
    from docx import Document

    path = _root() / "statistical_summary_table.docx"
    if not path.is_file():
        pytest.skip(str(path) + " does not exist; run scripts/35_statistical_reporting.py")
    assert path.stat().st_size > 10_000

    document = Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Paired fold-level comparisons" in text
    assert "Friedman" in text
    assert "McNemar" in text
    assert "Bootstrap confidence intervals" in text
    assert "sample size" in text
    assert "not a diagnostic tool" in text
    assert len(document.tables) >= 4


def test_t28_exists_in_every_format_and_states_what_its_n_means() -> None:
    import json

    root = _root()
    slug = "T28_statistical_significance_comparison"
    csv = root / (slug + ".csv")
    if not csv.is_file():
        pytest.skip(str(csv) + " does not exist; run scripts/35_statistical_reporting.py")
    for suffix in (".docx", ".tex", ".md", ".meta.json"):
        assert (root / (slug + suffix)).is_file(), slug + suffix

    import pandas as pd

    table = pd.read_csv(csv)
    assert {"n", "n_description", "p_value", "p_holm", "effect_size"} <= set(table.columns)
    assert table["n_description"].notna().all()

    notes = " ".join(
        json.loads((root / (slug + ".meta.json")).read_text(encoding="utf-8")).get("notes", [])
    )
    assert "EVERY p-VALUE STATES ITS n" in notes
    assert "EFFECT SIZE" in notes
    assert "Holm" in notes and "Benjamini" in notes


# ---------------------------------------------------------------------------
# clause 2 -- the correction is applied
# ---------------------------------------------------------------------------


def test_both_corrections_are_applied_and_are_never_smaller_than_the_raw_p() -> None:
    """A correction that lowered a p-value would be a sign the family is wrong."""
    paired = _csv("paired_fold_tests_interpreted.csv")
    for column in ("p_holm", "p_bh", "n_comparisons_in_family", "correction_family"):
        assert column in paired.columns, column

    reported = paired[paired["p_value"].notna()]
    assert reported["p_holm"].notna().all()
    assert reported["p_bh"].notna().all()
    assert (reported["p_holm"] >= reported["p_value"] - 1e-12).all()
    assert (reported["p_bh"] >= reported["p_value"] - 1e-12).all()
    # Holm is the more conservative of the two, always.
    assert (reported["p_holm"] >= reported["p_bh"] - 1e-12).all()
    assert (reported["n_comparisons_in_family"] > 1).any()


def test_the_correction_family_is_one_run_and_metric_not_everything_at_once() -> None:
    """Correcting across metrics would treat sensitivity and specificity on the
    same folds as independent tests, which they are not."""
    paired = _csv("paired_fold_tests_interpreted.csv")
    for family, block in paired.groupby("correction_family"):
        assert block["run"].nunique() == 1, family
        assert block["metric"].nunique() == 1, family


# ---------------------------------------------------------------------------
# clause 3 -- non-significant ensemble advantages are stated plainly
# ---------------------------------------------------------------------------


def test_the_findings_name_every_powered_ensemble_comparison_that_did_not_reject() -> None:
    from src.evaluation.statistics import ALPHA
    from src.reporting.statistics_report import ENSEMBLE_MODELS

    text = _findings()
    paired = _csv("paired_fold_tests_interpreted.csv")
    block = paired[paired["metric"] == "sensitivity"]
    involves = block[
        block["model_a"].isin(ENSEMBLE_MODELS) | block["model_b"].isin(ENSEMBLE_MODELS)
    ]
    ties = involves[
        (involves["p_holm"] >= ALPHA) & (~involves["underpowered"].astype(bool))
    ]
    assert len(ties), "expected at least one powered, non-significant ensemble comparison"
    for _, row in ties.iterrows():
        needle = str(row["model_a"]) + " and " + str(row["model_b"])
        assert needle in text, needle + " is not stated in the findings document"
    assert "indistinguishable" in text


def test_underpowered_and_non_significant_are_kept_apart() -> None:
    """They are different findings and the document must not merge them."""
    text = _findings()
    assert "## Where the proposed ensemble is NOT significantly better" in text
    assert "## Where the test was UNDERPOWERED rather than negative" in text
    assert "must never be reported as 'no difference'" in text

    tie_section = text.split("## Where the test was UNDERPOWERED")[0]
    assert "UNDERPOWERED, NOT NEGATIVE" not in tie_section, (
        "an underpowered comparison leaked into the negative-results section"
    )


def test_every_underpowered_run_is_named_with_its_floor() -> None:
    from src.evaluation.statistics import ALPHA

    text = _findings()
    paired = _csv("paired_fold_tests_interpreted.csv")
    weak = paired[
        (paired["metric"] == "sensitivity") & paired["underpowered"].astype(bool)
    ]
    for run in sorted(set(weak["run"])):
        assert run in text, run + " is underpowered but is not named"
    if len(weak):
        assert "0.0625" in text, "the n=5 floor must be stated as a number"
    assert str(ALPHA) in text


def test_the_findings_document_regenerates_identically(tmp_path) -> None:
    """Rule 5 applied to prose."""
    from src.reporting.statistics_report import write_findings

    first = _root() / "statistical_findings.md"
    if not first.is_file():
        pytest.skip(str(first) + " does not exist; run scripts/35_statistical_reporting.py")
    original = first.read_text(encoding="utf-8")

    import pandas as pd

    posthoc_path = _root() / "nemenyi_posthoc.csv"
    second = write_findings(
        tmp_path / "statistical_findings.md",
        paired=_csv("paired_fold_tests_interpreted.csv"),
        omnibus=_csv("friedman_omnibus.csv"),
        posthoc=pd.read_csv(posthoc_path) if posthoc_path.is_file() else None,
        sources=tuple(
            line.strip().strip("- `")
            for line in original.splitlines()
            if line.startswith("- `")
        ),
        command="python scripts/35_statistical_reporting.py",
    )
    assert second.read_text(encoding="utf-8") == original


def test_the_statistics_artifacts_are_registered_in_the_evidence_index() -> None:
    """T83.6 -- every artifact, not only the numbered table."""
    from src.utils.evidence import read_evidence

    rows = {row["evidence_id"]: row for row in read_evidence()}
    for evidence_id in ("STAT-01", "STAT-02", "STAT-03", "T28"):
        assert evidence_id in rows, evidence_id + " is not in the evidence index"
        assert rows[evidence_id]["status"] == "ok", evidence_id
        assert rows[evidence_id]["command"], evidence_id + " has no command"
