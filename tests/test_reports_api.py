"""Phase 117's API half: the report endpoints (T117.3) and the explanation (T117.2).

The reports page triggers three documents and the explainability page renders a
per-recording decomposition. Both come from the running API, so both are tested
here against the real generators, the real committed files and -- where the
checkout has them -- the real saved model and a real corpus recording.

What is pinned:

* An experiment report is a DOCX, carries the disclaimer, and names only runs the
  discovery found; an id cannot be a path.
* The objective-coverage report is T29's committed DOCX, byte for byte.
* A sample report needs a source and says so; with a real recording it is a DOCX
  carrying the disclaimer.
* The explanation the API returns is `feature_contributions` over the vector that
  was scored, term for term, formatted by `format_value`.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient

from src.api.main import EXPLANATION_TOP, create_app
from src.inference.predictor import DISCLAIMER

PROJECT_ROOT = Path(__file__).resolve().parents[1]
BUNDLE_DIR = PROJECT_ROOT / "models_saved" / "binary" / "final"
PHYSIONET = PROJECT_ROOT / "dataset" / "archive (3)" / "training-a"


@pytest.fixture(scope="module")
def client() -> Any:
    app = create_app(static_root=PROJECT_ROOT / "does-not-exist", preload=False)
    with TestClient(app) as running:
        yield running


def _docx_text(content: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


# ---------------------------------------------------------------------------
# experiment and objective reports -- committed files only, so they run on CI
# ---------------------------------------------------------------------------


def test_the_reportable_runs_are_discovered_not_typed() -> None:
    from src.reporting.pages_10_12 import reportable_experiments

    runs = reportable_experiments()
    assert "EXP-A1" in runs and "EXP-A2" in runs
    for exp_id, directory in runs.items():
        assert directory.name == exp_id
        assert (directory / "aggregate_metrics.csv").is_file()


def test_an_experiment_report_is_a_docx_carrying_the_disclaimer(client: Any) -> None:
    response = client.get("/report/experiment/EXP-A1")
    assert response.status_code == 200, response.text
    assert response.content[:2] == b"PK", "not a DOCX (zip) container"
    assert "EXP-A1" in response.headers["content-disposition"]
    text = _docx_text(response.content)
    assert DISCLAIMER in text
    assert "experiment report" in text


@pytest.mark.parametrize("exp_id", ["EXP-NOPE", "..", "..%2F..%2Fconfigs"])
def test_an_unknown_or_path_like_experiment_id_is_a_404(client: Any, exp_id: str) -> None:
    response = client.get("/report/experiment/" + exp_id)
    assert response.status_code == 404


def test_the_objective_report_is_t29s_committed_docx(client: Any) -> None:
    from src.reporting.pages_10_12 import objective_report_path

    path = objective_report_path()
    if path is None:
        pytest.skip("T29 has not been generated")
    response = client.get("/report/objectives")
    assert response.status_code == 200
    assert response.content == path.read_bytes()


def test_a_sample_report_without_a_source_is_a_400(client: Any) -> None:
    response = client.post("/report/sample", data={"task": "binary"})
    assert response.status_code == 400
    assert "sample_id or a file" in response.json()["detail"]


def test_a_sample_report_for_an_unknown_task_is_a_400(client: Any) -> None:
    response = client.post("/report/sample", data={"task": "merged", "sample_id": "x"})
    assert response.status_code == 400
    assert "never merged" in response.json()["detail"]


def test_an_empty_upload_is_refused_and_leaves_no_working_directory(client: Any) -> None:
    import tempfile

    spool = Path(tempfile.gettempdir())
    before = set(spool.glob("pvmepcg_report_*"))
    response = client.post(
        "/report/sample",
        data={"task": "binary"},
        files={"file": ("empty.wav", b"", "audio/wav")},
    )
    assert response.status_code == 400
    assert set(spool.glob("pvmepcg_report_*")) == before


# ---------------------------------------------------------------------------
# the real model on a real recording
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def recording() -> Path:
    if not (BUNDLE_DIR / "model.joblib").is_file():
        pytest.skip("no saved binary model in this checkout (gitignored)")
    if not PHYSIONET.is_dir():
        pytest.skip("dataset/ is not present in this checkout")
    recordings = sorted(PHYSIONET.glob("*.wav"))
    if not recordings:
        pytest.skip("no PhysioNet training-a recordings on disk")
    return recordings[0]


def test_a_sample_report_from_an_upload_is_a_docx_with_the_disclaimer(
    client: Any, recording: Path
) -> None:
    response = client.post(
        "/report/sample",
        data={"task": "binary"},
        files={"file": (recording.name, recording.read_bytes(), "audio/wav")},
    )
    assert response.status_code == 200, response.text
    text = _docx_text(response.content)
    assert DISCLAIMER in text
    assert "What drove this prediction" in text


def test_the_explanation_is_the_decomposition_of_the_scored_vector(
    client: Any, recording: Path
) -> None:
    from src.inference.predictor import predict_recording
    from src.reporting.sample_report import feature_contributions
    from src.reporting.tables import format_value

    served = client.post(
        "/predict",
        files={"file": (recording.name, recording.read_bytes(), "audio/wav")},
        data={"task": "binary"},
    ).json()
    explanation = served["explanation"]
    assert explanation["available"] is True, explanation
    assert len(explanation["rows"]) == EXPLANATION_TOP

    _, detail = predict_recording(recording, task="binary", with_detail=True)
    direct, reason = feature_contributions(detail.bundle, detail.vector, top=EXPLANATION_TOP)
    assert reason is None
    for shown, expected in zip(explanation["rows"], direct, strict=True):
        assert shown["feature"] == expected.name
        assert shown["contribution"] == expected.contribution
        assert shown["contribution_display"] == format_value(expected.contribution, "metric")
        assert shown["direction"] == expected.direction
