"""The evidence index, assembled and checked against both source documents (Phase 102).

Every generated artifact registers one row in ``evidence_index.csv`` as it is
written (``src/utils/evidence.py``). This module is the step that turns those
rows into a deliverable and asks whether they are enough:

T102.1/T102.2
    ``evidence_index.xlsx`` -- every row, the eleven spec fields, plus a sheet of
    mandatory items and a summary.
T102.3
    Every row is checked against the filesystem. The workbook's ``status`` uses
    the extraction spec's vocabulary: ``verified`` (the file exists and every
    source path it names exists), ``generated`` (the file exists but a source
    could not be resolved, or none was recorded), ``failed`` (the file is not
    on disk).
T102.4
    :data:`MANDATORY` is the list of every item the extraction spec
    (``KP_KHARAT_22366.docx``) and the developer blueprint name as required,
    each mapped to the evidence rows or paths that answer it. An item is
    ``present``, ``not produced`` (with a declared technical reason), or
    ``failed`` -- missing with no reason, which is the case the gate exists for.
T102.5
    :func:`finalize_manifest` adds a ``final`` block to ``run_manifest.json``:
    environment, packages, seed discipline across every run, and wall time per
    pipeline stage.
T102.6
    :func:`update_missing_report` regenerates one delimited block at the end of
    ``missing_outputs_report.txt`` listing every unproduced mandatory item. The
    hand-appended entries above it are never touched.

The catalogue is a literal, like the feature registry: which items the source
documents require is a fact about two documents, not something to derive.
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from src.utils.evidence import (
    EVIDENCE_COLUMNS,
    PROJECT_ROOT,
    _write_rows,
    evidence_index_path,
    read_evidence,
)
from src.utils.logging_setup import get_logger

__all__ = [
    "MANDATORY",
    "MandatoryItem",
    "BACKFILL",
    "WORKBOOK_NAME",
    "backfill_unregistered",
    "check_mandatory",
    "finalize_manifest",
    "merge_sidecars",
    "row_status",
    "update_missing_report",
    "write_evidence_workbook",
    "assemble",
]

log = get_logger(__name__)

WORKBOOK_NAME = "evidence_index.xlsx"
COMMAND = "python scripts/42_evidence_index.py"

SPEC = "extraction spec (KP_KHARAT_22366.docx)"
BLUEPRINT = "developer blueprint (PDF)"

#: A source token that names a path rather than describing one ("feature registry").
_PATHLIKE = re.compile(r"^(outputs|dataset|src|cache|models_saved|configs|scripts|Docs)/")

REPORT_BEGIN = (
    "=== BEGIN GENERATED COMPLETENESS AUDIT (Phase 102; rewritten by "
    "python scripts/42_evidence_index.py -- do not edit inside this block) ==="
)
REPORT_END = "=== END GENERATED COMPLETENESS AUDIT ==="


# ---------------------------------------------------------------------------
# declared reasons for items that are not produced
# ---------------------------------------------------------------------------

PART_X_SCREENSHOTS = (
    "Dashboard screenshots are Phase 120 (Part X). T119.4 makes the displayed-value "
    "audit a hard gate before any screenshot is taken, so a missing screenshot here "
    "means the gated capture (python scripts/47_dashboard_screenshots.py) has not been "
    "run in this checkout -- it needs a built frontend/out/, a running inference "
    "service and a local dataset/ copy."
)
PHASE_103 = "Produced by Phase 103 (Q1 / IEEE paper asset pack); regenerate this index after it."
PHASE_104 = "Produced by Phase 104 (thesis asset pack); regenerate this index after it."
ONE_COMMAND = (
    "The single command reproducing the pipeline end to end is "
    "scripts/00_run_everything.py (T122.1). Every artifact also carries its own "
    "reproduction command in this index, and tests/test_run_everything.py asserts "
    "that every script named in this column appears in that runner's stage list."
)
FINAL_ZIP = (
    "The delivery ZIP is built by `python scripts/51_package_delivery.py` into "
    "`dist/PV-MEPCG_PulseVision_delivery.zip` (125.5 MB, 2,228 files) and verified by "
    "reopening it -- CRCs, required entries, no excluded path, and a real "
    "frontend/out/index.html (T126.3/T126.4/T126.7). It is deliberately NOT given a "
    "path here: `dist/` is gitignored, and an item whose status depended on a local "
    "build product would make this committed report say one thing on the machine that "
    "built it and another on a fresh clone. The archive's own evidence row is "
    "DELIVERY-ZIP, written when it is built."
)
M9_REASON = (
    "M9 (1D-CNN) is out of scope, decided 2026-08-27 by the user: this machine has no GPU "
    "and the same 25-fold map is required for comparability (T52.4). See the M9 entry "
    "above in this file."
)
SAMPLE_OUTPUTS = (
    "No committed artifact holds per-recording system outputs. src/reporting/"
    "sample_report.py (Phase 107) generates a per-recording report on demand from an "
    "uploaded WAV, and its batch CSV/JSON export runs through the dashboard (Part X). "
    "A table of sample outputs would be generated through that path, not typed; the "
    "exported-report preview is dashboard screenshot 13 (Phase 120)."
)


# ---------------------------------------------------------------------------
# the catalogue
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class MandatoryItem:
    """One required deliverable and what in the repository answers it."""

    item_id: str
    source: str
    section: str
    description: str
    evidence_ids: tuple[str, ...] = ()
    paths: tuple[str, ...] = ()
    #: Declared technical reason the item does not exist yet. Consulted only
    #: when the item is absent: an item that exists is present regardless.
    reason: str | None = None
    note: str = ""


def _ids(prefix: str, count: int, width: int = 2) -> list[str]:
    return [prefix + str(n).zfill(width) for n in range(1, count + 1)]


SPEC_TABLE_NAMES = (
    "Dataset Inventory",
    "Class Distribution and Imbalance Ratio",
    "Recording Duration and Sampling Summary",
    "Preprocessing Configuration",
    "Feature Inventory and Counts",
    "Model Hyperparameter Configuration",
    "Search Space and Best Parameters",
    "PhysioNet Individual Model Comparison",
    "Equal-Weight vs Optimized Ensemble",
    "PhysioNet Fold-wise Results Mean +/- SD",
    "PASCAL A Multiclass Results",
    "PASCAL B Multiclass Results",
    "CirCor Murmur Results",
    "CirCor Clinical Outcome Results",
    "Recording-Level vs Subject-Level Results",
    "Cross-Dataset Generalization",
    "Feature-Family Ablation",
    "All Features vs Selected Features",
    "Search/Optimization Ablation",
    "Noise Robustness",
    "Duration Robustness",
    "Auscultation Location Analysis",
    "Calibration and Confidence Summary",
    "Complexity Analysis",
    "Training Time and Inference Time",
    "Model Size and Memory",
    "False Positive and False Negative Analysis",
    "Statistical Significance Comparison",
    "Objective-to-Evidence Mapping",
    "Final Conclusion Matrix",
)

SPEC_SCREENSHOTS = (
    "Home and project overview",
    "Dataset inventory and class distribution",
    "Signal upload and waveform",
    "Preprocessing before/after",
    "Feature extraction summary",
    "Model comparison dashboard",
    "Search optimization dashboard",
    "Binary prediction output",
    "Multiclass prediction output",
    "CirCor murmur/outcome output",
    "Robustness analytics",
    "Feature importance/explainability",
    "Exported report preview",
)

#: Section 15's twenty Q1 assets, by the id the Phase 103 pack registers.
Q1_ASSETS = (
    ("Q1_T01", "dataset summary"),
    ("Q1_T02", "feature summary"),
    ("Q1_T03", "model and search configuration"),
    ("Q1_T04", "individual vs ensemble"),
    ("Q1_T05", "feature ablation"),
    ("Q1_T06", "optimization ablation"),
    ("Q1_T07", "multiclass results"),
    ("Q1_T08", "external validation"),
    ("Q1_T09", "complexity and robustness"),
    ("Q1_F01", "proposed architecture"),
    ("Q1_F02", "feature and ensemble workflow"),
    ("Q1_G01", "ROC and PR combined"),
    ("Q1_G02", "binary confusion matrix"),
    ("Q1_G03", "multiclass confusion matrix"),
    ("Q1_G04", "ablation and optimization"),
    ("Q1_G05", "generalization and complexity"),
    ("Q1_ALG01", "feature extraction"),
    ("Q1_ALG02", "search optimization"),
    ("Q1_ALG03", "proposed ensemble"),
    ("Q1_NARRATIVE", "Q1_results_narrative.docx"),
)

THESIS_CHAPTERS = (
    "Ch1_Introduction",
    "Ch2_Literature",
    "Ch3_Methodology",
    "Ch4_Implementation",
    "Ch5_Results",
    "Ch6_Conclusion",
)

SPEC_FOLDERS = (
    "outputs/00_evidence_index",
    "outputs/01_dataset_audit",
    "outputs/02_preprocessing",
    "outputs/03_features",
    "outputs/04_models",
    "outputs/05_search_optimization",
    "outputs/06_binary_results",
    "outputs/07_multiclass_results",
    "outputs/08_circor_external_validation",
    "outputs/09_ablation",
    "outputs/10_robustness",
    "outputs/11_complexity",
    "outputs/12_statistics",
    "outputs/13_figures_diagrams",
    "outputs/14_algorithms",
    "outputs/15_dashboard_screenshots",
    "outputs/Q1_PAPER_ASSETS",
    "outputs/THESIS_ASSETS",
    "outputs/logs",
    "models_saved",
    "outputs/configs",
)


def _spec_items() -> list[MandatoryItem]:
    items: list[MandatoryItem] = []

    def add(item_id: str, section: str, description: str, **kwargs: Any) -> None:
        items.append(MandatoryItem(item_id, SPEC, section, description, **kwargs))

    for evidence_id in _ids("DA-", 9):
        add("SPEC-" + evidence_id, "3 dataset audit", evidence_id, evidence_ids=(evidence_id,))
    for evidence_id in _ids("PP-", 9):
        add("SPEC-" + evidence_id, "4 preprocessing", evidence_id, evidence_ids=(evidence_id,))
    for evidence_id in _ids("FE-", 12):
        add("SPEC-" + evidence_id, "5 features", evidence_id, evidence_ids=(evidence_id,))

    models = {
        "M1": "Logistic Regression (mandatory)",
        "M2": "K-Nearest Neighbours (optional)",
        "M3": "SVM-RBF (mandatory)",
        "M4": "Random Forest (mandatory)",
        "M5": "Gradient Boosting (mandatory)",
        "M6": "Soft voting, equal weight (mandatory)",
        "M7": "Soft voting, optimized weight (mandatory)",
        "M8": "XGBoost (recommended)",
    }
    for model_id, description in models.items():
        add("SPEC-" + model_id, "6 models", description, evidence_ids=("MODEL-" + model_id,))
    add("SPEC-M9", "6 models", "1D-CNN (optional)", reason=M9_REASON)

    searches = {
        "SO-01": ("SO-01",),
        "SO-02": ("SO-02",),
        "SO-03": ("SO-03a", "SO-03b"),
        "SO-04": ("SO-04",),
        "SO-05": ("SO-05",),
        "SO-06": ("SO-06",),
    }
    for item, ids in searches.items():
        add("SPEC-" + item, "7 search and optimization", item, evidence_ids=ids)

    runs = {
        "EXP-A1": ("EXP-A1", "EXP-A1-AGG"),
        "EXP-A2": ("EXP-A2", "EXP-A2-AGG"),
        "EXP-B1": ("EXP-B1", "EXP-B1-AGG"),
        "EXP-B2": ("EXP-B2", "EXP-B2-AGG"),
        "EXP-C1": ("EXP-C1-three_class", "EXP-C1-three_class-AGG", "EXP-C1-two_class"),
        "EXP-C2": ("EXP-C2", "EXP-C2-AGG"),
        "EXP-C3": ("T22", "G32"),
        "EXP-D1": ("T16", "G29"),
        "EXP-E1": ("T20", "G30"),
        "EXP-E2": ("T21", "G31"),
        "EXP-F1": (*("EXP-F1-A" + str(n) for n in range(1, 9)), "T17"),
        "EXP-F2": ("T19", "EXP-F2-A9-DEFAULT", "EXP-F2-A9-TUNED", "EXP-F2-A10"),
    }
    for item, ids in runs.items():
        add("SPEC-" + item, "8 experimental runs", item, evidence_ids=ids)

    for number, name in enumerate(SPEC_TABLE_NAMES, start=1):
        table_id = "T" + str(number).zfill(2)
        add("SPEC-" + table_id, "9 tables", name, evidence_ids=(table_id,))
    for graph_id in _ids("G", 35):
        add("SPEC-" + graph_id, "10 graphs", graph_id, evidence_ids=(graph_id,))
    for figure_id in _ids("F", 20):
        add("SPEC-" + figure_id, "11 figures and diagrams", figure_id, evidence_ids=(figure_id,))
    for alg_id in _ids("ALG-", 20):
        add("SPEC-" + alg_id, "12 algorithms", alg_id, evidence_ids=(alg_id,))

    add(
        "SPEC-STAT-MATRIX",
        "13 statistics",
        "statistical_significance_matrix.csv",
        evidence_ids=("STAT-01",),
    )
    add(
        "SPEC-STAT-SUMMARY",
        "13 statistics",
        "statistical_summary_table.docx",
        evidence_ids=("STAT-02",),
    )
    add(
        "SPEC-STAT-BOOTSTRAP",
        "13 statistics",
        "bootstrap ROC-AUC confidence intervals for key binary results",
        evidence_ids=("STAT-BOOTSTRAP",),
    )

    for number, name in enumerate(SPEC_SCREENSHOTS, start=1):
        add(
            "SPEC-SS" + str(number).zfill(2),
            "14 dashboard screenshots",
            name,
            evidence_ids=("SS-" + str(number).zfill(2),),
            reason=PART_X_SCREENSHOTS,
        )

    for evidence_id, name in Q1_ASSETS:
        add(
            "SPEC-" + evidence_id,
            "15 Q1 paper assets",
            name,
            evidence_ids=(evidence_id,),
            reason=PHASE_103,
        )

    for chapter in THESIS_CHAPTERS:
        add(
            "SPEC-THESIS-" + chapter,
            "16 thesis assets",
            "chapter folder " + chapter,
            evidence_ids=("THESIS-" + chapter,),
            reason=PHASE_104,
        )
    add(
        "SPEC-THESIS-REPRO",
        "16 thesis assets",
        "complete reproducibility appendix",
        evidence_ids=("THESIS-REPRODUCIBILITY",),
        reason=PHASE_104,
    )
    add(
        "SPEC-THESIS-MATRICES",
        "16 thesis assets",
        "objective achievement matrix and final conclusion matrix",
        evidence_ids=("T29", "T30"),
    )
    add(
        "SPEC-THESIS-SCREENSHOTS",
        "16 thesis assets",
        "dashboard screenshots with captions",
        evidence_ids=tuple("SS-" + str(n).zfill(2) for n in range(1, 14)),
        reason=PART_X_SCREENSHOTS,
    )

    for folder in SPEC_FOLDERS:
        add("SPEC-DIR-" + folder.split("/")[-1], "17 folder structure", folder, paths=(folder,))
    add(
        "SPEC-EVIDENCE-XLSX",
        "18 evidence index",
        "outputs/00_evidence_index/evidence_index.xlsx",
        paths=("outputs/00_evidence_index/" + WORKBOOK_NAME,),
    )
    add(
        "SPEC-RUN-MANIFEST",
        "17 folder structure",
        "outputs/00_evidence_index/run_manifest.json",
        paths=("outputs/00_evidence_index/run_manifest.json",),
    )
    add("SPEC-README", "19 checklist", "README", paths=("README.md",))
    add(
        "SPEC-ONE-COMMAND",
        "19 checklist",
        "one-command reproduction instructions",
        paths=("scripts/00_run_everything.py",),
        reason=ONE_COMMAND,
    )
    add(
        "SPEC-MISSING-REPORT",
        "20 final delivery",
        "missing_outputs_report.txt",
        paths=("outputs/missing_outputs_report.txt",),
    )
    add("SPEC-ZIP", "20 final delivery", "final delivery ZIP", reason=FINAL_ZIP)
    return items


#: Blueprint section 9's twenty tables, numbered as the blueprint numbers them.
BLUEPRINT_TABLES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("Dataset Inventory", ("T01",)),
    ("Class Distribution", ("T02",)),
    ("Preprocessing Settings", ("T04",)),
    ("Feature Inventory", ("T05",)),
    ("Model Configuration", ("T06",)),
    ("Search Space", ("T07",)),
    ("PhysioNet Binary Comparison", ("T08",)),
    ("PASCAL Dataset A Results", ("T11",)),
    ("PASCAL Dataset B Results", ("T12",)),
    ("CirCor External Results", ("T13", "T14", "T15")),
    ("Cross-Dataset Generalization", ("T16",)),
    ("Feature Ablation", ("T17",)),
    ("Optimization Ablation", ("T19",)),
    ("Complexity Analysis", ("T24", "T25", "T26")),
    ("Noise Robustness", ("T20",)),
    ("Duration Robustness", ("T21",)),
    ("Calibration and Confidence", ("T23",)),
    ("Failure Analysis", ("T27",)),
    ("Objective Achievement", ("T29",)),
    ("Final System Outputs", ()),
)

#: Blueprint section 10's G1-G20, mapped onto the extraction spec's G01-G35.
BLUEPRINT_GRAPHS: tuple[tuple[str, str], ...] = (
    ("Dataset class-distribution bar chart", "G02"),
    ("Waveform before/after filtering", "G05"),
    ("Spectrogram comparison", "G06"),
    ("Feature-family count chart", "G10"),
    ("Model comparison bar chart", "G11"),
    ("ROC curve", "G12"),
    ("Precision-recall curve", "G13"),
    ("Binary confusion matrix", "G14"),
    ("Multiclass confusion matrix A", "G15"),
    ("Multiclass confusion matrix B", "G16"),
    ("Feature-importance plot", "G18"),
    ("Search convergence plot", "G20"),
    ("Accuracy/F1 vs feature count", "G22"),
    ("Inference time vs performance", "G25"),
    ("Dataset-wise generalization chart", "G28"),
    ("Noise-level robustness chart", "G30"),
    ("Duration-wise performance chart", "G31"),
    ("Confidence histogram", "G33"),
    ("Calibration curve", "G34"),
    ("Error-analysis chart", "G35"),
)

BLUEPRINT_ABLATIONS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("A1 time-domain only", ("EXP-F1-A1",)),
    ("A2 frequency-domain only", ("EXP-F1-A2",)),
    ("A3 MFCC only", ("EXP-F1-A3",)),
    ("A4 wavelet only", ("EXP-F1-A4",)),
    ("A5 time + frequency", ("EXP-F1-A5",)),
    ("A6 MFCC + wavelet", ("EXP-F1-A6",)),
    ("A7 all 138 features", ("EXP-F1-A7",)),
    ("A8 optimized feature subset", ("EXP-F1-A8",)),
    ("A9 best individual model vs ensemble", ("EXP-F2-A9-DEFAULT", "EXP-F2-A9-TUNED")),
    ("A10 equal-weight vs optimized-weight ensemble", ("EXP-F2-A10", "T09")),
)


def _blueprint_items() -> list[MandatoryItem]:
    items: list[MandatoryItem] = []

    def add(item_id: str, section: str, description: str, **kwargs: Any) -> None:
        items.append(MandatoryItem(item_id, BLUEPRINT, section, description, **kwargs))

    for number in range(1, 7):
        add(
            "BP-OBJ" + str(number),
            "1/3 locked objectives",
            "Objective " + str(number) + " evidence",
            evidence_ids=("T29", "T30"),
        )
    add(
        "BP-OBJ2-LIT",
        "3 objective mapping",
        "Objective 2: literature review table",
        evidence_ids=("LIT-01", "LIT-02"),
    )
    add(
        "BP-OBJ2-ABL",
        "3 objective mapping",
        "Objective 2: preprocessing ablation and feature-family ablation",
        evidence_ids=("PP-09", "T17"),
    )

    for model_id, name in (
        ("M1", "Logistic Regression"),
        ("M3", "SVM-RBF"),
        ("M4", "Random Forest"),
        ("M5", "Gradient Boosting"),
        ("M6", "Soft Voting Ensemble"),
    ):
        add("BP-MODEL-" + model_id, "7 algorithm plan", name, evidence_ids=("MODEL-" + model_id,))
    add(
        "BP-MODEL-M2",
        "7 algorithm plan",
        "K-Nearest Neighbours (optional)",
        evidence_ids=("MODEL-M2",),
    )
    add(
        "BP-MODEL-M8",
        "7 algorithm plan",
        "XGBoost or LightGBM (recommended)",
        evidence_ids=("MODEL-M8",),
    )
    add("BP-MODEL-M9", "7 algorithm plan", "1D-CNN (optional)", reason=M9_REASON)

    add(
        "BP-SEARCH-GRID",
        "7 search-method plan",
        "Grid search on a small parameter space",
        evidence_ids=("SO-05",),
        note=(
            "M7's ensemble weights are chosen by an exhaustive deterministic grid over the "
            "weight simplex (232 candidates); SO-05 is that search."
        ),
    )
    add("BP-SEARCH-RANDOM", "7 search-method plan", "RandomizedSearchCV", evidence_ids=("SO-01",))
    add("BP-SEARCH-BAYES", "7 search-method plan", "Bayesian optimization", evidence_ids=("SO-02",))
    add("BP-SEARCH-GA", "7 search-method plan", "Genetic algorithm", evidence_ids=("SO-03a",))
    add(
        "BP-SEARCH-PSO",
        "7 search-method plan",
        "Particle swarm optimization",
        evidence_ids=("SO-03b",),
    )

    tracks = {
        "A": ("Track A PhysioNet binary", ("EXP-A1", "EXP-A2")),
        "B": ("Track B PASCAL A four-class", ("EXP-B1",)),
        "C": ("Track C PASCAL B three-class", ("EXP-B2",)),
        "D": ("Track D CirCor murmur and outcome", ("EXP-C1-three_class", "EXP-C2")),
        "E": ("Track E cross-dataset", ("T16",)),
        "F": ("Track F ablation", ("T17", "T19")),
    }
    for letter, (name, ids) in tracks.items():
        add("BP-TRACK-" + letter, "8 experimental tracks", name, evidence_ids=ids)
    for index, (name, ids) in enumerate(BLUEPRINT_ABLATIONS, start=1):
        add("BP-ABL-A" + str(index), "8 ablation plan", name, evidence_ids=ids)

    for number, (name, ids) in enumerate(BLUEPRINT_TABLES, start=1):
        add(
            "BP-TABLE" + str(number).zfill(2),
            "9 tables",
            "Table " + str(number) + " " + name,
            evidence_ids=ids,
            reason=SAMPLE_OUTPUTS if not ids else None,
        )
    for number, (name, graph_id) in enumerate(BLUEPRINT_GRAPHS, start=1):
        add(
            "BP-G" + str(number).zfill(2),
            "10 graphs",
            "G" + str(number) + " " + name,
            evidence_ids=(graph_id,),
        )

    add(
        "BP-EQUATIONS",
        "11 equations",
        "equations and optimization formulas",
        evidence_ids=("EQUATIONS-CSV", "EQUATIONS-TEX", "EQUATIONS-DOCX"),
    )
    add("BP-BINARY-CM", "12 binary module", "binary confusion matrix", evidence_ids=("G14",))
    add(
        "BP-MULTICLASS",
        "13 multiclass module",
        "macro-F1, weighted-F1, per-class recall, OvR AUC, confusion matrices",
        evidence_ids=("T11", "T12", "G15", "G16", "G17"),
    )
    add(
        "BP-EXTERNAL",
        "14 external validation",
        "recording- and subject-level aggregation, location analysis, failure report",
        evidence_ids=("T15", "T22", "T27"),
    )
    add(
        "BP-DASHBOARD",
        "15 dashboard",
        "dashboard pages, evidenced by the section 14 screenshots",
        evidence_ids=tuple("SS-" + str(n).zfill(2) for n in range(1, 14)),
        reason=PART_X_SCREENSHOTS,
    )
    for evidence_id, name in Q1_ASSETS:
        if evidence_id.startswith(("Q1_T", "Q1_G", "Q1_NARRATIVE")):
            add(
                "BP-" + evidence_id,
                "16 Q1 result package",
                name,
                evidence_ids=(evidence_id,),
                reason=PHASE_103,
            )
    for chapter in THESIS_CHAPTERS:
        add(
            "BP-" + chapter,
            "17 thesis chapters",
            chapter,
            evidence_ids=("THESIS-" + chapter,),
            reason=PHASE_104,
        )
    add(
        "BP-SAMPLE-REPORTS",
        "19 final deliverables",
        "sample-level exportable reports",
        reason=SAMPLE_OUTPUTS,
    )
    add(
        "BP-README",
        "19 final deliverables",
        "README and requirements",
        paths=("README.md", "requirements"),
    )
    add(
        "BP-SAVED-MODELS",
        "19 final deliverables",
        "saved baseline, tuned and ensemble models",
        evidence_ids=(
            "MODEL-M1",
            "MODEL-M3",
            "MODEL-M4",
            "MODEL-M5",
            "MODEL-M6",
            "MODEL-M7",
            "T65-FINAL-MODEL",
        ),
    )
    return items


MANDATORY: tuple[MandatoryItem, ...] = tuple(_spec_items() + _blueprint_items())


# ---------------------------------------------------------------------------
# T102.1: artifacts generated before the index existed, or never registered
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Backfill:
    evidence_id: str
    filename: str
    metric_or_asset: str
    source_data: str
    command: str
    objective: str = ""
    experiment_id: str = ""
    dataset: str = ""
    model: str = ""


_SMOKE = "python scripts/04_model_smoke.py"
_MATRIX = "outputs/03_features/all_features_matrix.parquet"
_FIGURES = 'python -c "from src.feature_extraction.figures import generate_all; generate_all()"'
_ABLATION = "python scripts/27_optimization_ablation.py"

BACKFILL: tuple[Backfill, ...] = (
    Backfill(
        "FE-01",
        "outputs/03_features/feature_inventory.csv",
        "all 138 features with family, extractor, equation reference and unit",
        "src/feature_extraction/registry.py",
        'python -c "from src.feature_extraction.extractor import write_feature_artifacts; '
        'write_feature_artifacts()"',
        objective="O4 (feature engineering)",
    ),
    Backfill(
        "FE-02",
        "outputs/03_features/feature_family_summary.csv",
        "six feature families and their locked counts",
        "src/feature_extraction/registry.py",
        'python -c "from src.feature_extraction.extractor import write_feature_artifacts; '
        'write_feature_artifacts()"',
        objective="O4 (feature engineering)",
    ),
    Backfill(
        "FE-07",
        "outputs/03_features/chroma_heatmap.png",
        "chroma heatmap, representative record",
        "outputs/01_dataset_audit/metadata_master.csv",
        _FIGURES,
        objective="O4 (feature engineering)",
        dataset="D1",
    ),
    Backfill(
        "FE-08",
        "outputs/03_features/wavelet_decomposition.png",
        "db4 multi-level decomposition, representative record",
        "outputs/01_dataset_audit/metadata_master.csv",
        _FIGURES,
        objective="O4 (feature engineering)",
        dataset="D1",
    ),
    Backfill(
        "FE-09",
        "outputs/03_features/signal_envelope.png",
        "amplitude envelope and detected peaks, representative record",
        "outputs/01_dataset_audit/metadata_master.csv",
        _FIGURES,
        objective="O4 (feature engineering)",
        dataset="D1",
    ),
    *(
        Backfill(
            "MODEL-" + model_id,
            "models_saved/binary/" + model_id + "/manifest.json",
            "saved " + model_id + " (binary, fold 0 smoke fit): size, timing, features, versions",
            _MATRIX,
            _SMOKE,
            objective="O3 (model training)",
            dataset="D1 PhysioNet 2016",
            model=model_id,
        )
        for model_id in ("M1", "M2", "M3", "M4", "M5", "M6", "M7", "M8")
    ),
    Backfill(
        "EXP-F2-A9-DEFAULT",
        "outputs/09_ablation/A9_individual_vs_ensemble_default.csv",
        "A9 best individual model vs ensemble, default parameters",
        "outputs/09_ablation/optimization_ablation_per_fold.csv",
        _ABLATION,
        objective="O3 (optimization)",
        experiment_id="EXP-F2",
        dataset="D1 PhysioNet 2016",
    ),
    Backfill(
        "EXP-F2-A9-TUNED",
        "outputs/09_ablation/A9_individual_vs_ensemble_tuned.csv",
        "A9 best individual model vs ensemble, tuned parameters",
        "outputs/09_ablation/optimization_ablation_per_fold.csv",
        _ABLATION,
        objective="O3 (optimization)",
        experiment_id="EXP-F2",
        dataset="D1 PhysioNet 2016",
    ),
    Backfill(
        "EXP-F2-A10",
        "outputs/09_ablation/A10_equal_vs_optimized_weights.csv",
        "A10 equal-weight vs optimized-weight ensemble",
        "outputs/09_ablation/optimization_ablation_per_fold.csv",
        _ABLATION,
        objective="O3 (optimization)",
        experiment_id="EXP-F2",
        dataset="D1 PhysioNet 2016",
    ),
    Backfill(
        "STAT-BOOTSTRAP",
        "outputs/12_statistics/bootstrap_auc_ci.csv",
        "bootstrap ROC-AUC confidence intervals, key binary results",
        "; ".join(
            directory + "/predictions.parquet"
            for directory in (
                "outputs/06_binary_results/EXP-A1",
                "outputs/06_binary_results/EXP-A2",
                "outputs/07_multiclass_results/EXP-B1",
                "outputs/07_multiclass_results/EXP-B2",
                "outputs/08_circor_external_validation/EXP-C1-three_class",
                "outputs/08_circor_external_validation/EXP-C1-two_class",
                "outputs/08_circor_external_validation/EXP-C2",
            )
        ),
        "python scripts/34_statistical_validation.py",
        objective="O3 (model comparison)",
        experiment_id="EXP-A1, EXP-A2",
        dataset="D1 PhysioNet 2016",
    ),
)


def _mtime_utc(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime, tz=UTC).isoformat()


def backfill_unregistered(index_path: str | Path | None = None) -> list[str]:
    """Register the :data:`BACKFILL` artifacts that have no row yet.

    Each row's timestamp is the file's modification time, not now: the field
    means "generation time", and these files were generated before the index
    looked at them. A row that already exists is left alone -- the phase that
    wrote it knows more than this list does.
    """
    target = Path(index_path) if index_path else evidence_index_path()
    rows = read_evidence(target)
    have = {row["evidence_id"] for row in rows}
    added: list[str] = []
    for item in BACKFILL:
        if item.evidence_id in have:
            continue
        artifact = PROJECT_ROOT / item.filename
        rows.append(
            {
                "evidence_id": item.evidence_id,
                "objective": item.objective,
                "experiment_id": item.experiment_id,
                "dataset": item.dataset,
                "model": item.model,
                "metric_or_asset": item.metric_or_asset,
                "filename": item.filename,
                "source_data": item.source_data,
                "command": item.command,
                "timestamp": _mtime_utc(artifact) if artifact.is_file() else "",
                "status": "ok" if artifact.is_file() else "missing",
            }
        )
        added.append(item.evidence_id)
    if added:
        _write_rows(rows, target)
    return added


# ---------------------------------------------------------------------------
# the parallel-session sidecars (2026-08-30 entry in note.md)
# ---------------------------------------------------------------------------


def merge_sidecars(directory: str | Path | None = None, *, remove: bool = True) -> dict[str, Any]:
    """Fold ``evidence_index_window2.csv`` and ``run_manifest_window2.json`` in.

    Evidence upserts on ``evidence_id`` -- but a row the main index already
    has wins, because every id present in both was re-registered in the main
    index after the parallel session ended. Runs concatenate on ``run_id``;
    config snapshots merge by digest (content-hashed, so a collision is
    identical by construction).
    """
    folder = Path(directory) if directory else evidence_index_path().parent
    evidence_side = folder / "evidence_index_window2.csv"
    manifest_side = folder / "run_manifest_window2.json"
    report: dict[str, Any] = {"evidence_added": [], "evidence_kept_main": [], "runs_added": 0}

    if evidence_side.is_file():
        main_path = folder / "evidence_index.csv"
        rows = read_evidence(main_path)
        have = {row["evidence_id"] for row in rows}
        for row in read_evidence(evidence_side):
            if row["evidence_id"] in have:
                report["evidence_kept_main"].append(row["evidence_id"])
            else:
                rows.append(row)
                report["evidence_added"].append(row["evidence_id"])
        _write_rows(rows, main_path)
        if remove:
            evidence_side.unlink()

    if manifest_side.is_file():
        from src.utils.io import save_json

        main_manifest = folder / "run_manifest.json"
        data = json.loads(main_manifest.read_text(encoding="utf-8"))
        side = json.loads(manifest_side.read_text(encoding="utf-8"))
        known = {run["run_id"] for run in data["runs"]}
        added = [run for run in side.get("runs", []) if run["run_id"] not in known]
        data["runs"] = sorted(data["runs"] + added, key=lambda run: run.get("started_utc") or "")
        for digest, snapshot in side.get("config_snapshots", {}).items():
            data["config_snapshots"].setdefault(digest, snapshot)
        save_json(data, main_manifest)
        report["runs_added"] = len(added)
        if remove:
            manifest_side.unlink()
    return report


# ---------------------------------------------------------------------------
# T102.3: each row against the filesystem
# ---------------------------------------------------------------------------


def _resolve(path: str) -> Path:
    candidate = Path(path)
    return candidate if candidate.is_absolute() else PROJECT_ROOT / candidate


def source_paths(source_data: str) -> list[str]:
    """The tokens of a ``source_data`` field that name a path."""
    tokens = [token.strip() for token in re.split(r";", source_data or "") if token.strip()]
    # "dataset/ (read-only input)": the path is the part before the gloss.
    tokens = [re.sub(r"\s+\(.*\)$", "", token) for token in tokens]
    return [token for token in tokens if _PATHLIKE.match(token)]


def row_status(row: dict[str, str]) -> tuple[str, str]:
    """``(status, detail)`` in the spec's vocabulary: verified / generated / failed."""
    if not row.get("filename") or not _resolve(row["filename"]).is_file():
        return "failed", "file not on disk: " + (row.get("filename") or "(no filename)")
    sources = source_paths(row.get("source_data", ""))
    if not row.get("source_data"):
        return "generated", "no source_data recorded"
    unresolved = [source for source in sources if not _resolve(source).exists()]
    if unresolved:
        return "generated", "source not on disk: " + "; ".join(unresolved)
    if not sources:
        return "verified", "file on disk; source is descriptive, not a path"
    return "verified", "file and " + str(len(sources)) + " source path(s) on disk"


# ---------------------------------------------------------------------------
# T102.4: the mandatory items
# ---------------------------------------------------------------------------


def ignored_by_git(path: str) -> bool:
    """True when git ignores ``path`` -- a file a fresh clone can never have."""
    import subprocess

    result = subprocess.run(
        ["git", "check-ignore", "-q", path],
        cwd=PROJECT_ROOT,
        capture_output=True,
        check=False,
    )
    return result.returncode == 0


def fails_only_on_ignored_files(item: dict[str, str], rows: list[dict[str, str]]) -> bool:
    """True when a failed item fails ONLY because its files are gitignored.

    The standing rule: a check needing a gitignored input must skip, not fail.
    FE-03 (the feature matrix parquet) is the case -- present locally, absent on
    every CI runner by design. An id missing from the index is never excused.
    """
    by_id = {row["evidence_id"]: row for row in rows}
    ids = [evidence_id for evidence_id in item["evidence_ids"].split("; ") if evidence_id]
    if any(evidence_id not in by_id for evidence_id in ids):
        return False
    missing = [
        by_id[evidence_id]["filename"]
        for evidence_id in ids
        if row_status(by_id[evidence_id])[0] == "failed"
    ]
    missing += [path for path in item["paths"].split("; ") if path and not _resolve(path).exists()]
    return bool(missing) and all(ignored_by_git(path) for path in missing)


def check_mandatory(rows: list[dict[str, str]] | None = None) -> list[dict[str, str]]:
    """One result per :data:`MANDATORY` item: present / not produced / failed."""
    rows = read_evidence() if rows is None else rows
    by_id = {row["evidence_id"]: row for row in rows}
    results: list[dict[str, str]] = []
    for item in MANDATORY:
        problems: list[str] = []
        for evidence_id in item.evidence_ids:
            row = by_id.get(evidence_id)
            if row is None:
                problems.append(evidence_id + " not in the index")
            elif row_status(row)[0] == "failed":
                problems.append(evidence_id + " file missing")
        for path in item.paths:
            if not _resolve(path).exists():
                problems.append(path + " not on disk")
        answered = bool(item.evidence_ids or item.paths)
        if answered and not problems:
            status, detail = "present", ""
        elif item.reason:
            status, detail = "not produced", item.reason
        else:
            status, detail = "failed", "; ".join(problems)
        results.append(
            {
                "item_id": item.item_id,
                "source_document": item.source,
                "section": item.section,
                "description": item.description,
                "evidence_ids": "; ".join(item.evidence_ids),
                "paths": "; ".join(item.paths),
                "status": status,
                "detail": detail,
                "note": item.note,
            }
        )
    return results


# ---------------------------------------------------------------------------
# T102.1/T102.2: the workbook
# ---------------------------------------------------------------------------

_MANDATORY_COLUMNS = (
    "item_id",
    "source_document",
    "section",
    "description",
    "evidence_ids",
    "paths",
    "status",
    "detail",
    "note",
)


def _sheet(
    workbook: Any, title: str, columns: tuple[str, ...] | list[str], rows: list[dict[str, str]]
) -> None:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(title)
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(columns, start=1):
        longest = max([len(str(column))] + [len(str(row.get(column, ""))) for row in rows[:500]])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 70)


def write_evidence_workbook(
    target: str | Path,
    rows: list[dict[str, str]],
    mandatory: list[dict[str, str]],
) -> Path:
    """Three sheets: every row, every mandatory item, and the counts."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)

    indexed = []
    for row in rows:
        status, detail = row_status(row)
        indexed.append({**row, "status": status, "status_detail": detail})
    _sheet(workbook, "evidence_index", [*EVIDENCE_COLUMNS, "status_detail"], indexed)
    _sheet(workbook, "mandatory_items", _MANDATORY_COLUMNS, mandatory)

    row_counts = Counter(row["status"] for row in indexed)
    item_counts = Counter(item["status"] for item in mandatory)
    summary = [
        {"measure": "generated_utc", "value": datetime.now(UTC).isoformat()},
        {"measure": "command", "value": COMMAND},
        {"measure": "framework", "value": "PV-MEPCG / PulseVision"},
        {"measure": "evidence rows", "value": str(len(indexed))},
        *(
            {"measure": "rows " + key, "value": str(row_counts.get(key, 0))}
            for key in ("verified", "generated", "failed")
        ),
        {"measure": "mandatory items", "value": str(len(mandatory))},
        *(
            {"measure": "items " + key, "value": str(item_counts.get(key, 0))}
            for key in ("present", "not produced", "failed")
        ),
        {
            "measure": "status vocabulary",
            "value": (
                "verified = file and every named source path on disk; generated = file on "
                "disk, a source unresolved or unrecorded; failed = file not on disk"
            ),
        },
        {
            "measure": "disclaimer",
            "value": (
                "Academic screening and decision-support prototype. Not a diagnostic tool; "
                "no diagnosis, treatment or prescription is made or implied."
            ),
        },
    ]
    _sheet(workbook, "summary", ("measure", "value"), summary)

    path = Path(target)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(path))
    return path


# ---------------------------------------------------------------------------
# T102.5: the manifest's final block
# ---------------------------------------------------------------------------


def _seconds(run: dict[str, Any]) -> float | None:
    total = (run.get("timings_sec") or {}).get("_total_run")
    if total is not None:
        return float(total)
    if run.get("started_utc") and run.get("finished_utc"):
        started = datetime.fromisoformat(run["started_utc"])
        finished = datetime.fromisoformat(run["finished_utc"])
        return (finished - started).total_seconds()
    return None


def finalize_manifest(path: str | Path | None = None) -> dict[str, Any]:
    """Add (or refresh) the manifest's ``final`` block. Runs are never altered.

    Runs record the pipeline stage (the script's run name), not a todo.md phase
    number, so timing is reported per stage. A run still marked ``running`` is a
    process that ended without calling ``finish()`` -- killed, crashed or
    interrupted -- and is counted, never given a made-up duration.
    """
    from src.utils.io import save_json
    from src.utils.run_manifest import (
        _default_manifest_path,
        environment_info,
        git_info,
        package_versions,
    )
    from src.utils.seed import GLOBAL_SEED

    target = Path(path) if path else _default_manifest_path()
    data = json.loads(target.read_text(encoding="utf-8"))
    runs = data.get("runs", [])

    stages: dict[str, dict[str, Any]] = {}
    for run in runs:
        stage = stages.setdefault(
            run.get("name", "run"),
            {
                "stage": run.get("name", "run"),
                "n_runs": 0,
                "n_finished": 0,
                "n_unfinished": 0,
                "total_wall_seconds": 0.0,
                "max_wall_seconds": 0.0,
                "first_started_utc": None,
                "last_finished_utc": None,
            },
        )
        stage["n_runs"] += 1
        seconds = _seconds(run)
        if run.get("finished_utc") and seconds is not None:
            stage["n_finished"] += 1
            stage["total_wall_seconds"] = round(stage["total_wall_seconds"] + seconds, 3)
            stage["max_wall_seconds"] = round(max(stage["max_wall_seconds"], seconds), 3)
            if (
                stage["last_finished_utc"] is None
                or run["finished_utc"] > stage["last_finished_utc"]
            ):
                stage["last_finished_utc"] = run["finished_utc"]
        else:
            stage["n_unfinished"] += 1
        started = run.get("started_utc")
        if started and (stage["first_started_utc"] is None or started < stage["first_started_utc"]):
            stage["first_started_utc"] = started

    seeds = Counter(run.get("seed") for run in runs)
    data["final"] = {
        "finalized_utc": datetime.now(UTC).isoformat(),
        "finalized_by": COMMAND,
        "task": "T102.5",
        "framework": "PV-MEPCG / PulseVision",
        "environment": environment_info(),
        "packages": package_versions(),
        "git": git_info(),
        "seed": {
            "global_seed": GLOBAL_SEED,
            "runs_by_seed": {str(key): value for key, value in seeds.items()},
            "runs_with_another_seed": [
                run["run_id"] for run in runs if run.get("seed") != GLOBAL_SEED
            ],
        },
        "runs": {
            "total": len(runs),
            "by_status": dict(Counter(run.get("status", "unknown") for run in runs)),
            "unfinished_note": (
                "status 'running' means the process ended without finish() -- killed, "
                "crashed or interrupted; its duration is not recorded and not estimated"
            ),
        },
        "config_snapshots": len(data.get("config_snapshots", {})),
        "stage_timing": sorted(stages.values(), key=lambda stage: stage["first_started_utc"] or ""),
    }
    save_json(data, target)
    return data["final"]


# ---------------------------------------------------------------------------
# T102.6: the generated block of missing_outputs_report.txt
# ---------------------------------------------------------------------------


def missing_report_path() -> Path:
    return PROJECT_ROOT / "outputs" / "missing_outputs_report.txt"


def update_missing_report(mandatory: list[dict[str, str]], path: str | Path | None = None) -> Path:
    """Rewrite the generated block; leave every hand-appended entry untouched."""
    target = Path(path) if path else missing_report_path()
    text = target.read_text(encoding="utf-8") if target.is_file() else ""
    if REPORT_BEGIN in text:
        head, _, rest = text.partition(REPORT_BEGIN)
        _, _, tail = rest.partition(REPORT_END)
        text = head.rstrip("\n") + "\n" + tail.lstrip("\n")
    text = text.rstrip("\n") + "\n\n"

    absent = [item for item in mandatory if item["status"] != "present"]
    today = datetime.now(UTC).date().isoformat()
    lines = [
        REPORT_BEGIN,
        "Generated " + datetime.now(UTC).isoformat() + ". Every mandatory item of the",
        "extraction spec and the developer blueprint that is not present on disk, with its",
        "technical reason. "
        + str(len(mandatory))
        + " items checked; "
        + str(sum(1 for i in mandatory if i["status"] == "present"))
        + " present; "
        + str(len(absent))
        + " listed below. The full table is the 'mandatory_items' sheet",
        "of outputs/00_evidence_index/evidence_index.xlsx.",
        "",
    ]
    for item in absent:
        lines.append(
            "["
            + today
            + "] "
            + item["item_id"]
            + " -- "
            + item["description"]
            + " ("
            + item["source_document"]
            + ", section "
            + item["section"]
            + ")"
        )
        lines.append("  Status: " + item["status"])
        label = "Reason: " if item["status"] == "not produced" else "Failure: "
        lines.append("  " + label + item["detail"])
        lines.append("")
    if not absent:
        lines.append("(none)")
        lines.append("")
    lines.append(REPORT_END)
    target.write_text(text + "\n".join(lines) + "\n", encoding="utf-8")
    return target


# ---------------------------------------------------------------------------
# one call
# ---------------------------------------------------------------------------


def assemble(*, index_path: str | Path | None = None) -> dict[str, Any]:
    """The whole of Phase 102, in the order the tasks need it."""
    merged = merge_sidecars()
    added = backfill_unregistered(index_path)
    rows = read_evidence(index_path)
    mandatory = check_mandatory(rows)

    folder = Path(index_path).parent if index_path else evidence_index_path().parent
    workbook = write_evidence_workbook(folder / WORKBOOK_NAME, rows, mandatory)
    # The workbook is itself a mandatory item: check again now that it exists.
    mandatory = check_mandatory(rows)
    workbook = write_evidence_workbook(folder / WORKBOOK_NAME, rows, mandatory)

    final = finalize_manifest()
    report = update_missing_report(mandatory)
    statuses = Counter(row_status(row)[0] for row in rows)
    return {
        "merged": merged,
        "backfilled": added,
        "workbook": workbook,
        "report": report,
        "rows": len(rows),
        "row_status": dict(statuses),
        "items": dict(Counter(item["status"] for item in mandatory)),
        "failed_items": [item for item in mandatory if item["status"] == "failed"],
        "failed_rows": [row for row in rows if row_status(row)[0] == "failed"],
        "stages": len(final["stage_timing"]),
    }


def read_workbook_rows(path: str | Path, sheet: str) -> list[dict[str, str]]:
    """A sheet as dicts -- what a gate reads back."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True)
    values = list(workbook[sheet].iter_rows(values_only=True))
    header = [str(cell) for cell in values[0]]
    return [
        {header[i]: ("" if cell is None else str(cell)) for i, cell in enumerate(row)}
        for row in values[1:]
    ]


def read_csv_rows(path: str | Path) -> list[dict[str, str]]:
    with Path(path).open("r", encoding="utf-8", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]
