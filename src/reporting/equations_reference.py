"""The equations reference (Phase 100): section 11 as LaTeX, DOCX and CSV, tied to file:line.

:mod:`src.reporting.equations` declares the fifteen formulas of blueprint
section 11 once -- their LaTeX, their symbol tables, the module that implements
each -- and the dashboard renders them from there. This module renders the same
declarations for the paper and the thesis, and adds the one thing T100.5 asks
for that the dashboard does not need: **the line of code that computes each
formula**.

The line is found, not typed. Each equation names a fragment of its implementing
code (:data:`CODE_REFERENCES`); the file is searched for it at export time and
the matching line number is what the reference prints. A formula whose code has
been rewritten therefore fails the export instead of citing a line that now
does something else -- the same contract ``_quoted`` gives the ALG pseudocode.

Three outputs, all in ``outputs/14_algorithms/``:

* ``equations_reference.tex`` -- a standalone ``amsmath`` document; each formula
  is a numbered, labelled ``equation`` ready to paste into the paper (T100.3);
* ``equations_reference.docx`` -- every formula and every symbol as a native
  Office Math object, editable in Word's equation editor (T100.4, T100.6),
  converted from the same LaTeX by :mod:`src.reporting.omml`;
* ``equations_reference.csv`` -- one row per equation with its file, line and
  code, the machine-readable cross-reference.

Where the implementation computes a documented formula in a specific form -- a
normalised entropy, a band-pass rather than a low-pass response -- the
difference is stated in :data:`IMPLEMENTATION_NOTES` rather than left for a
reader to discover. ``tests/test_equations_reference.py`` checks every formula
numerically against the code (T100.2, T100.7).
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from src.reporting.equations import EQUATIONS, Equation
from src.utils.logging_setup import get_logger

__all__ = [
    "CODE_REFERENCES",
    "IMPLEMENTATION_NOTES",
    "CrossReference",
    "cross_reference",
    "reference_rows",
    "write_equations_reference",
]

log = get_logger("reporting.equations_reference")

COMMAND = "python scripts/40_export_equations.py"
STEM = "equations_reference"

#: ``key -> fragment`` of the line that computes the formula, in the equation's
#: ``implemented_in`` file. Searched for at export; a missing fragment is an error.
CODE_REFERENCES: dict[str, str] = {
    "butterworth": "magnitude = 1.0 / np.sqrt(1.0 + transformed ** (2 * order))",
    "zscore": "((samples - samples.mean()) / std)",
    "energy": "energy = float(np.sum(signal**2))",
    "rms": "rms = float(np.sqrt(energy / signal.size))",
    "spectral_centroid": (
        "centroid = librosa.feature.spectral_centroid(S=magnitude, sr=fs, n_fft=n_fft)"
    ),
    "spectral_entropy": "entropy = float(-np.sum(probabilities * np.log2(probabilities)))",
    "soft_voting": "return np.tensordot(vector / total, stack, axes=(0, 0))",
    "prediction": "return classes[np.argmax(proba, axis=1)]",
    "accuracy": '"accuracy": float((tp + tn) / (tp + tn + fp + fn)),',
    "sensitivity": "recall_score(true, pred, pos_label=positive_label, zero_division=0)",
    "specificity": 'return float(tn / denominator) if denominator else float("nan")',
    "f1": "f1_score(true, pred, pos_label=positive_label, zero_division=0)",
    "balanced_accuracy": '"balanced_accuracy": float(np.nanmean([sensitivity, specificity])),',
    "macro_f1": 'kwargs: dict[str, Any] = {"average": "macro", "zero_division": 0}',
    "objective_j": "weights.alpha * (1.0 - float(performance))",
}

#: How the implementation instantiates a documented formula, where that is not
#: the formula read literally. Each is verified numerically by the tests.
IMPLEMENTATION_NOTES: dict[str, str] = {
    "butterworth": (
        "The equation is the low-pass prototype. The implemented filter is a band-pass, "
        "reached by the substitution W = (w^2 - w0^2) / (w BW) with w0^2 = w1 w2 and "
        "BW = w2 - w1, after bilinear prewarping of every frequency; the reference "
        "function reproduces scipy's sosfreqz of the designed filter. The filter is run "
        "forward and backward, so the applied magnitude is the square of this response."
    ),
    "energy": (
        "Computed on the z-normalized signal, where it equals the number of samples; the "
        "time-domain feature time_energy is therefore duration x sampling rate."
    ),
    "rms": "Computed on the z-normalized signal, where it is identically 1 (time_rms).",
    "spectral_centroid": (
        "Computed per short-time Fourier frame by librosa, with X_k the frame's magnitude "
        "spectrum, and summarised as the mean and SD over frames (freq_centroid_mean, "
        "freq_centroid_std)."
    ),
    "spectral_entropy": (
        "Implemented with base-2 logarithms on the Welch power spectrum and divided by "
        "log2 of the number of bins, so the feature lies in [0, 1] and is comparable "
        "between records whose spectra have different lengths."
    ),
    "soft_voting": (
        "Weights are normalised by their sum before averaging, so raw scores may be "
        "passed; M6 uses equal weights and M7 the weights chosen in-fold (ALG-11, ALG-12)."
    ),
    "prediction": (
        "The argmax rule is the multiclass decision of SoftVotingEnsemble.predict. For the "
        "binary task the ensembles predict by a threshold chosen inside the training fold, "
        "and the deployed model predicts by argmax at the fixed 0.5 reference."
    ),
    "macro_f1": (
        "Defined for the binary task too, where it averages the F1 of both classes -- not "
        "the positive-class F1 of equation 12."
    ),
    "objective_j": (
        "SelectedFeatures is the subset size and 138 the declared total "
        "(JWeights.n_features_total). NormalizedInferenceTime is the measured extraction "
        "cost of the feature families the subset still needs, divided by the cost of all "
        "families."
    ),
}


class EquationReferenceError(RuntimeError):
    """An equation cannot be cross-referenced to the code that computes it."""


@dataclass(frozen=True)
class CrossReference:
    path: str
    line: int
    code: str

    def cited(self) -> str:
        return self.path + ":" + str(self.line)


def _root() -> Path:
    return Path(__file__).resolve().parents[2]


def _normal(text: str) -> str:
    return " ".join(text.split())


def cross_reference(equation: Equation) -> CrossReference:
    """``(file, line, code)`` of the line that computes ``equation``. Raises if absent."""
    fragment = CODE_REFERENCES.get(equation.key)
    if not fragment:
        raise EquationReferenceError("equation " + equation.key + " names no code fragment")
    path = _root() / equation.implemented_in
    if not path.is_file():
        raise EquationReferenceError(equation.implemented_in + " does not exist")
    lines = path.read_text(encoding="utf-8").splitlines()
    wanted = _normal(fragment)
    for number, line in enumerate(lines, start=1):
        if wanted in _normal(line):
            return CrossReference(equation.implemented_in, number, line.strip())
    raise EquationReferenceError(
        "equation "
        + str(equation.number)
        + " ("
        + equation.key
        + ") cites "
        + repr(fragment)
        + ", which no longer occurs in "
        + equation.implemented_in
    )


def reference_rows() -> list[dict[str, Any]]:
    """One row per equation, with its resolved cross-reference."""
    rows: list[dict[str, Any]] = []
    for equation in EQUATIONS:
        reference = cross_reference(equation)
        rows.append(
            {
                "number": equation.number,
                "key": equation.key,
                "name": equation.name,
                "latex": equation.latex,
                "use": equation.use,
                "implemented_in": reference.path,
                "line": reference.line,
                "code": reference.code,
                "implements": equation.implements,
                "symbols": "; ".join(
                    symbol + " = " + meaning for symbol, meaning in equation.symbols
                ),
                "transcription_note": equation.transcription_note or "",
                "implementation_note": IMPLEMENTATION_NOTES.get(equation.key, ""),
            }
        )
    return rows


# ---------------------------------------------------------------------------
# writers
# ---------------------------------------------------------------------------


def _write_csv(rows: list[dict[str, Any]], target: Path) -> Path:
    with target.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    return target


def _tex_text(text: str) -> str:
    """Escape prose for LaTeX, leaving ``$...$`` spans as math."""
    from src.reporting.tables import _latex_escape

    parts = text.split("$")
    return "".join(
        ("$" + part + "$") if index % 2 else _latex_escape(part) for index, part in enumerate(parts)
    )


def _write_latex(rows: list[dict[str, Any]], target: Path) -> Path:
    from src.reporting.result_tables import DISCLAIMER

    by_key = {equation.key: equation for equation in EQUATIONS}
    out = [
        "% PV-MEPCG / PulseVision -- equations of blueprint section 11.",
        "% Generated by " + COMMAND + " from src/reporting/equations.py. Do not edit;",
        "% every formula is cross-referenced to the line of code that computes it.",
        r"\documentclass{article}",
        r"\usepackage{amsmath}",
        r"\usepackage{booktabs}",
        r"\begin{document}",
        r"\section*{PV-MEPCG / PulseVision: equations and optimization formulas}",
        _tex_text(
            "The fifteen formulas of the developer blueprint, section 11, with their "
            "symbols and the source line that implements each. " + DISCLAIMER
        ),
    ]
    for row in rows:
        equation = by_key[str(row["key"])]
        out += [
            "",
            r"\subsection*{" + _tex_text(str(row["name"])) + "}",
            "% implemented in " + str(row["implemented_in"]) + ":" + str(row["line"]),
            "% " + str(row["code"]),
            r"\begin{equation}\label{eq:" + str(row["key"]) + "}",
            equation.latex,
            r"\end{equation}",
            r"\noindent\textit{Use:} " + _tex_text(str(row["use"])) + r".\par",
            r"\smallskip\noindent\begin{tabular}{@{}ll@{}}",
            r"\toprule Symbol & Meaning \\ \midrule",
        ]
        out += [
            "$" + symbol + "$ & " + _tex_text(meaning) + r" \\"
            for symbol, meaning in equation.symbols
        ]
        out += [
            r"\bottomrule",
            r"\end{tabular}\par",
            r"\smallskip\noindent\textit{Implemented in} \texttt{"
            + _tex_text(str(row["implemented_in"]) + ":" + str(row["line"]))
            + r"}.\par",
        ]
        for label, key in (
            ("Transcription", "transcription_note"),
            ("Implementation", "implementation_note"),
        ):
            if row[key]:
                out.append(
                    r"\noindent\textit{" + label + " note.} " + _tex_text(str(row[key])) + r"\par"
                )
    out += ["", r"\end{document}", ""]
    target.write_text("\n".join(out), encoding="utf-8")
    return target


def _add_math(paragraph: Any, latex: str, *, display: bool) -> None:
    from docx.oxml import parse_xml

    from src.reporting.omml import latex_to_omml

    paragraph._p.append(parse_xml(latex_to_omml(latex, display=display)))


def _add_mixed(paragraph: Any, text: str, size: Any) -> None:
    """Prose with ``$...$`` spans rendered as inline Office Math."""
    for index, part in enumerate(text.split("$")):
        if index % 2:
            _add_math(paragraph, part, display=False)
        elif part:
            paragraph.add_run(part).font.size = size


def _write_docx(rows: list[dict[str, Any]], target: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    from src.reporting.result_tables import DISCLAIMER

    by_key = {equation.key: equation for equation in EQUATIONS}
    document = Document()
    title = document.add_paragraph()
    run = title.add_run("PV-MEPCG / PulseVision — Equations and optimization formulas")
    run.bold = True
    run.font.size = Pt(13)
    intro = document.add_paragraph(
        "The fifteen formulas of the developer blueprint, section 11. Every formula and "
        "symbol below is a native equation object, editable in Word's equation editor, "
        "and each names the source line that implements it."
    )
    intro.runs[0].font.size = Pt(9)

    for row in rows:
        equation = by_key[str(row["key"])]
        heading = document.add_paragraph()
        head = heading.add_run("(" + str(row["number"]) + ") " + str(row["name"]))
        head.bold = True
        head.font.size = Pt(11)
        _add_math(document.add_paragraph(), equation.latex, display=True)

        use = document.add_paragraph()
        key = use.add_run("Use: ")
        key.bold = True
        key.font.size = Pt(9)
        use.add_run(str(row["use"])).font.size = Pt(9)

        grid = document.add_table(rows=1, cols=2)
        grid.style = "Table Grid"
        for cell, header in zip(grid.rows[0].cells, ("Symbol", "Meaning"), strict=True):
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.size = Pt(8)
        for symbol, meaning in equation.symbols:
            cells = grid.add_row().cells
            _add_math(cells[0].paragraphs[0], symbol, display=False)
            _add_mixed(cells[1].paragraphs[0], meaning, Pt(8))

        cited = document.add_paragraph()
        label = cited.add_run("Implemented in: ")
        label.font.size = Pt(8)
        where = cited.add_run(str(row["implemented_in"]) + ":" + str(row["line"]))
        where.font.name = "Consolas"
        where.font.size = Pt(8)
        code = document.add_paragraph(str(row["code"]))
        code.runs[0].font.name = "Consolas"
        code.runs[0].font.size = Pt(7.5)
        for name in ("transcription_note", "implementation_note"):
            if row[name]:
                note = document.add_paragraph(
                    (
                        "Transcription note: "
                        if name == "transcription_note"
                        else "Implementation note: "
                    )
                    + str(row[name])
                )
                note.runs[0].italic = True
                note.runs[0].font.size = Pt(8)

    footer = document.add_paragraph(
        DISCLAIMER + "  |  generated from src/reporting/equations.py  |  regenerate: " + COMMAND
    )
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(7)
    document.save(str(target))
    return target


def write_equations_reference(
    out_dir: str | Path | None = None, *, evidence_index: str | Path | None = None
) -> dict[str, Path]:
    """Write the CSV, LaTeX and DOCX references and register them as evidence."""
    from src.reporting.algorithms import algorithms_dir
    from src.utils.evidence import register_evidence
    from src.utils.io import ensure_dir

    rows = reference_rows()
    target = Path(ensure_dir(algorithms_dir(out_dir)))
    written = {
        "csv": _write_csv(rows, target / (STEM + ".csv")),
        "tex": _write_latex(rows, target / (STEM + ".tex")),
        "docx": _write_docx(rows, target / (STEM + ".docx")),
    }
    sources = "; ".join(
        ["src/reporting/equations.py", *sorted({str(row["implemented_in"]) for row in rows})]
    )
    for kind, path in written.items():
        register_evidence(
            "EQUATIONS-" + kind.upper(),
            path,
            metric_or_asset="Equations reference, blueprint section 11 (" + kind + ")",
            objective="O1-O6",
            source_data=sources,
            command=COMMAND,
            index_path=evidence_index,
        )
    log.info("equations reference: %s", ", ".join(p.name for p in written.values()))
    return written
