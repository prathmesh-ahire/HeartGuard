"""T28, the significance matrix and the statistical summary (Phase 83).

Four artifacts over the Phase 82 tests, and each exists because a different
reader needs a different shape:

``statistical_significance_matrix.csv`` (T83.1)
    The literal model-by-model p-value matrix, one square block per
    (run, metric, test, correction). What a reader with a specific pair in mind
    wants.

``statistical_summary_table.docx`` (T83.2)
    Test, statistic, p-value, effect size and a written interpretation, one row
    per comparison. What goes into a thesis chapter.

``T28`` (T83.3)
    The numbered comparison table, through the ordinary table engine so it
    renders in four formats and carries its provenance like every other table.

``statistical_findings.md`` (T83.5)
    The plain statements. **Where the proposed ensemble is not significantly
    better than a simpler model, that is written down**, and so is the separate
    case where a test was underpowered rather than negative. Those two are
    different findings and the document keeps them apart.

## The interpretation column is generated, not written

Every sentence in it is assembled from the row: the chosen test, its n, the
corrected p-value, the effect size and whether the sample size could have
detected anything. A hand-written interpretation would drift from its numbers
the first time the analysis was re-run.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import numpy as np

from src.evaluation.statistics import ALPHA
from src.reporting.tables import Column, Table, TableSpec, build_table
from src.utils.io import ensure_dir
from src.utils.logging_setup import get_logger

__all__ = [
    "ENSEMBLE_MODELS",
    "significance_matrix",
    "interpret",
    "build_t28",
    "write_summary_docx",
    "write_findings",
]

log = get_logger("reporting.statistics")

#: The models the blueprint calls "the proposed ensemble". T83.5 asks
#: specifically whether these beat the simpler alternatives.
ENSEMBLE_MODELS: tuple[str, ...] = ("M6", "M7")

_DISCLAIMER = (
    "PV-MEPCG / PulseVision is an academic screening and decision-support "
    "prototype, not a diagnostic tool."
)


def significance_matrix(paired: Any, *, value: str = "p_holm") -> tuple[Any, Any]:
    """T83.1 -- one square p-value block per (run, metric), plus which test made it.

    Returns ``(p_values, tests)``: two frames of the same shape, the first
    holding the p-value in each cell and the second the name of the test that
    produced it.

    **Grouped by (run, metric) only, deliberately.** The first version also
    grouped by the chosen test, and the gate caught what that does: the test is
    chosen *per pair* by a normality check, so M1-M3 (paired t) and M3-M4
    (Wilcoxon) land in different blocks and each block is a mostly-empty
    triangle that is not symmetric. A reader looking up one pair would find a
    blank and have no way to know a value existed in the other block. One matrix
    per (run, metric) is the deliverable T83.1 asks for; the companion frame
    keeps the per-cell test visible rather than dropping it.

    Symmetric by construction: a paired test's p-value does not depend on which
    model is called A. The diagonal is left empty rather than filled with 1.0 --
    a model compared with itself was not tested, and a 1.0 there would read as a
    test that found no difference.
    """
    import pandas as pd

    p_blocks: list[pd.DataFrame] = []
    test_blocks: list[pd.DataFrame] = []
    keys = [k for k in ("run", "metric") if k in paired.columns]
    test_column = next((c for c in ("chosen_test", "test") if c in paired.columns), None)

    for values, block in paired.groupby(keys, sort=True):
        models = sorted(set(block["model_a"]) | set(block["model_b"]))
        matrix = pd.DataFrame(index=models, columns=models, dtype=float)
        which = pd.DataFrame(index=models, columns=models, dtype=object)
        for _, row in block.iterrows():
            p = float(row[value]) if value in row and pd.notna(row[value]) else float("nan")
            matrix.loc[row["model_a"], row["model_b"]] = p
            matrix.loc[row["model_b"], row["model_a"]] = p
            name = str(row[test_column]) if test_column else ""
            which.loc[row["model_a"], row["model_b"]] = name
            which.loc[row["model_b"], row["model_a"]] = name

        tests_used = "; ".join(sorted(set(block[test_column].astype(str)))) if test_column else ""
        for frame in (matrix, which):
            frame.reset_index(inplace=True)
            frame.rename(columns={"index": "model"}, inplace=True)
            for offset, (key, item) in enumerate(zip(keys, np.atleast_1d(values), strict=True)):
                frame.insert(offset, key, item)
            frame.insert(len(keys) + 1, "n", int(block["n"].iloc[0]))
            frame.insert(len(keys) + 2, "n_description", str(block["n_description"].iloc[0]))
            frame.insert(len(keys) + 3, "tests_used", tests_used)
        matrix.insert(len(keys) + 1, "p_value_kind", value)
        which.insert(len(keys) + 1, "cell_holds", "the test that produced the p-value")
        p_blocks.append(matrix)
        test_blocks.append(which)

    if not p_blocks:
        return pd.DataFrame(), pd.DataFrame()
    return (
        pd.concat(p_blocks, ignore_index=True),
        pd.concat(test_blocks, ignore_index=True),
    )


def interpret(row: Any) -> str:
    """One comparison in a sentence, assembled entirely from the row."""
    a, b = str(row["model_a"]), str(row["model_b"])
    n = int(row["n"])
    test = str(row.get("chosen_test") or row.get("test") or "test")
    raw = float(row.get("p_value", float("nan")))
    holm = float(row.get("p_holm", float("nan")))
    effect = float(row.get("effect_size", float("nan")))
    effect_name = str(row.get("effect_size_name", "effect size"))
    difference = float(row.get("mean_difference", float("nan")))
    underpowered = bool(row.get("underpowered", False))

    leader = a if difference > 0 else b if difference < 0 else "neither"
    size = abs(difference)

    if underpowered and not (holm < ALPHA):
        return (
            "UNDERPOWERED, NOT NEGATIVE. "
            + test
            + " on n="
            + str(n)
            + " cannot produce a p-value below "
            + format(float(row.get("minimum_achievable_p", float("nan"))), ".4f")
            + " however large the effect, so failing to reject says nothing about "
            + a
            + " versus "
            + b
            + ". The observed gap is "
            + format(size, ".4f")
            + " in favour of "
            + leader
            + "."
        )
    if not np.isfinite(holm):
        return (
            "No test result: the comparison could not be computed (see the "
            "p-value column). n=" + str(n) + "."
        )
    if holm < ALPHA:
        return (
            leader
            + " is significantly better ("
            + test
            + ", n="
            + str(n)
            + ", p="
            + format(raw, ".4g")
            + ", Holm-corrected p="
            + format(holm, ".4g")
            + "), by "
            + format(size, ".4f")
            + " with "
            + effect_name
            + " "
            + format(effect, ".3f")
            + "."
        )
    return (
        a
        + " and "
        + b
        + " are statistically indistinguishable ("
        + test
        + ", n="
        + str(n)
        + ", p="
        + format(raw, ".4g")
        + ", Holm-corrected p="
        + format(holm, ".4g")
        + "). The observed gap is "
        + format(size, ".4f")
        + " in favour of "
        + leader
        + ", with "
        + effect_name
        + " "
        + format(effect, ".3f")
        + " -- a difference this test had the power to detect and did not."
    )


def build_t28(
    paired: Any, sources: tuple[str, ...], command: str = "", metric: str = "sensitivity"
) -> Table:
    """T28 -- the numbered significance comparison, on the selection metric."""
    frame = paired[paired["metric"] == metric].copy()
    if not len(frame):
        frame = paired.copy()
    frame["interpretation"] = [interpret(row) for _, row in frame.iterrows()]
    frame = frame.sort_values(["run", "model_a", "model_b"]).reset_index(drop=True)

    columns = (
        Column("run", "Run"),
        Column("metric", "Metric"),
        Column("model_a", "Model A"),
        Column("model_b", "Model B"),
        Column("chosen_test", "Test"),
        Column("n", "n", kind="integer"),
        # T82.7 requires every reported p-value to state its n. "25" alone does
        # not: 25 folds and 25 records are different claims, and the gate that
        # caught this was right to.
        Column("n_description", "What n is"),
        Column("mean_a", "Mean A", kind="metric"),
        Column("mean_b", "Mean B", kind="metric"),
        Column("mean_difference", "Difference", kind="metric", places=4),
        Column("statistic", "Statistic", kind="metric", places=3),
        # p_value, not metric: at four places a p of 2.7e-14 printed "0.0000",
        # which reads as exactly zero. The p_value kind prints "<0.0001".
        Column("p_value", "p (raw)", kind="p_value", places=4),
        Column("p_holm", "p (Holm)", kind="p_value", places=4),
        Column("p_bh", "p (BH)", kind="p_value", places=4),
        Column("effect_size_name", "Effect size"),
        Column("effect_size", "Value", kind="metric", places=3),
        Column("underpowered", "Underpowered"),
        Column("interpretation", "Interpretation"),
    )
    present = tuple(c for c in columns if c.name in frame.columns)

    spec = TableSpec(
        table_id="T28",
        title="Statistical Significance Comparison",
        caption=(
            "Every model pair on "
            + metric
            + ", tested over fold-level paired observations. The test is chosen "
            "per pair by a Shapiro-Wilk check on the paired differences -- the "
            "paired t-test where they are normal, Wilcoxon signed-rank where they "
            "are not -- and both results are kept in the source CSV so the choice "
            "is auditable. p-values are reported raw and under two "
            "multiple-comparison corrections."
        ),
        sources=sources,
        columns=present,
        exp_id="EXP-A1, EXP-A2, EXP-B1, EXP-B2, EXP-C1, EXP-C2",
        objective="O3 (model comparison)",
        dataset="D1 PhysioNet 2016, D2/D3 PASCAL, D4 CirCor 2022",
        notes=(
            "EVERY p-VALUE STATES ITS n, AND n IS THE NUMBER OF FOLDS. The "
            "repeated 5x5 binary map gives n=25. EXP-B2 is a plain 5-fold map, so "
            "n=5 there, and a two-sided Wilcoxon on five pairs cannot go below "
            "0.0625 however large the effect. The 'Underpowered' column marks "
            "exactly those rows: a non-significant result there is a statement "
            "about the design, not about the models.",
            "AN EFFECT SIZE ACCOMPANIES EVERY p-VALUE. A significant result on 25 "
            "folds can sit on a 0.001 difference; the effect size is what stops "
            "that being written up as a finding.",
            "TWO CORRECTIONS, BOTH REPORTED. Holm controls the family-wise error "
            "rate and is the conservative choice for a claim about one pair; "
            "Benjamini-Hochberg controls the false discovery rate and is the "
            "right one for screening many pairs. The family is one (run, metric) "
            "group -- correcting across metrics would treat sensitivity and "
            "specificity on the same folds as independent tests.",
            "Fold-level tests use folds as the unit of observation. Record-level "
            "tests (McNemar, the AUC bootstrap) run on ONE repeat of the map, "
            "which is a complete partition of the corpus; pooling all 25 folds "
            "would put every record in the sample five times.",
            _DISCLAIMER,
        ),
        command=command,
    )
    return build_table(spec, frame)


def write_summary_docx(
    path: str | Path,
    *,
    paired: Any,
    omnibus: Any,
    mcnemar: Any,
    bootstrap: Any,
    command: str = "",
) -> Path:
    """T83.2 -- one DOCX row per test: what was run, on what, and what it means.

    Deliberately not routed through the table engine: this file is named in the
    task by its own filename rather than by a table id, and it carries four
    different kinds of test whose columns do not align into one grid.
    """
    from docx import Document
    from docx.shared import Pt

    document = Document()
    heading = document.add_paragraph()
    run_text = heading.add_run("Statistical summary")
    run_text.bold = True
    run_text.font.size = Pt(12)

    intro = document.add_paragraph(
        "Every test below states its sample size and what that sample size is. "
        "A fold-level test's n is the number of cross-validation folds; a "
        "record-level test's n is the number of records, each appearing exactly "
        "once. Generated by " + command + "."
    )
    intro.runs[0].font.size = Pt(9)

    sections: list[tuple[str, Any, tuple[str, ...]]] = [
        (
            "Paired fold-level comparisons (T82.4, T82.6)",
            paired,
            (
                "run",
                "metric",
                "model_a",
                "model_b",
                "chosen_test",
                "n",
                "statistic",
                "p_value",
                "p_holm",
                "effect_size_name",
                "effect_size",
                "interpretation",
            ),
        ),
        (
            "Friedman omnibus across all models (T82.5)",
            omnibus,
            (
                "run",
                "metric",
                "test",
                "n",
                "k_models",
                "statistic",
                "p_value",
                "effect_size_name",
                "effect_size",
                "average_ranks",
                "posthoc_note",
            ),
        ),
        (
            "McNemar on paired predictions (T82.3)",
            mcnemar,
            (
                "run",
                "test",
                "model_a",
                "model_b",
                "n",
                "statistic",
                "p_value",
                "p_holm",
                "effect_size_name",
                "effect_size",
                "favours",
            ),
        ),
        (
            "Bootstrap confidence intervals for ROC-AUC (T82.2)",
            bootstrap,
            (
                "run",
                "model_id",
                "metric",
                "n",
                "point",
                "ci_low",
                "ci_high",
                "n_resamples",
                "ci_kind",
            ),
        ),
    ]

    for title, frame, wanted in sections:
        section = document.add_paragraph()
        section_run = section.add_run(title)
        section_run.bold = True
        section_run.font.size = Pt(10)

        if frame is None or not len(frame):
            missing = document.add_paragraph(
                "Not produced. See outputs/missing_outputs_report.txt."
            )
            missing.runs[0].italic = True
            missing.runs[0].font.size = Pt(9)
            continue

        columns = [c for c in wanted if c in frame.columns]
        grid = document.add_table(rows=1, cols=len(columns))
        grid.style = "Table Grid"
        for cell, header in zip(grid.rows[0].cells, columns, strict=True):
            cell.text = header.replace("_", " ")
            for paragraph in cell.paragraphs:
                for cell_run in paragraph.runs:
                    cell_run.bold = True
                    cell_run.font.size = Pt(7)
        for _, row in frame.iterrows():
            cells = grid.add_row().cells
            for cell, column in zip(cells, columns, strict=True):
                value = row[column]
                cell.text = (
                    format(float(value), ".4g")
                    if isinstance(value, (int, float, np.floating)) and np.isfinite(value)
                    else str(value)
                )
                for paragraph in cell.paragraphs:
                    for cell_run in paragraph.runs:
                        cell_run.font.size = Pt(7)

    footer = document.add_paragraph(_DISCLAIMER)
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(8)

    target = Path(path)
    ensure_dir(target.parent)
    document.save(str(target))
    log.info("wrote %s", target)
    return target


def write_findings(
    path: str | Path,
    *,
    paired: Any,
    omnibus: Any,
    posthoc: Any,
    sources: tuple[str, ...],
    command: str = "",
    metric: str = "sensitivity",
) -> Path:
    """T83.5 -- the plain statements, including the ones nobody wants to write.

    Two sections that must not be merged: comparisons where the ensemble is
    **not significantly better**, and comparisons where the test **could not have
    detected a difference at all**. The first is a result about the models; the
    second is a result about the study design.
    """
    lines: list[str] = []
    lines.append("# Statistical findings")
    lines.append("")
    lines.append(
        "Generated by `" + command + "`. Every number is read from the frames "
        "named at the end. " + _DISCLAIMER
    )
    lines.append("")

    block = paired[paired["metric"] == metric]

    lines.append("## Where the proposed ensemble is NOT significantly better (T83.5)")
    lines.append("")
    lines.append(
        "Comparisons on **"
        + metric
        + "** involving "
        + " or ".join(ENSEMBLE_MODELS)
        + " where the Holm-corrected p-value does not reach "
        + str(ALPHA)
        + ", and the test had the power to detect a difference. These are real "
        "negative results, not missing ones."
    )
    lines.append("")
    involves = block[
        block["model_a"].isin(ENSEMBLE_MODELS) | block["model_b"].isin(ENSEMBLE_MODELS)
    ]
    ties = involves[(involves["p_holm"] >= ALPHA) & (~involves["underpowered"].astype(bool))]
    if len(ties):
        for _, row in ties.iterrows():
            lines.append("- **" + str(row["run"]) + "** -- " + interpret(row))
    else:
        lines.append(
            "- None: on every powered comparison at this metric the ensemble's "
            "difference from the other model reached significance."
        )
    lines.append("")

    lines.append("## Where the test was UNDERPOWERED rather than negative")
    lines.append("")
    lines.append(
        "A two-sided signed-rank test on n pairs has a smallest possible "
        "p-value of 2 x 0.5^n. Where that floor is at or above "
        + str(ALPHA)
        + ", the test could not have rejected however large the effect, and a "
        "non-significant result is a statement about the design. **These must "
        "never be reported as 'no difference'.**"
    )
    lines.append("")
    weak = block[block["underpowered"].astype(bool)]
    if len(weak):
        for run in sorted(set(weak["run"])):
            part = weak[weak["run"] == run]
            n = int(part["n"].iloc[0])
            floor = float(part["minimum_achievable_p"].iloc[0])
            lines.append(
                "- **"
                + run
                + "** -- n="
                + str(n)
                + " folds, smallest achievable p-value "
                + format(floor, ".4f")
                + ". "
                + str(len(part))
                + " comparison(s) affected; none of them can reach "
                + str(ALPHA)
                + "."
            )
    else:
        lines.append(
            "- None at this metric: every run tested here has enough folds for "
            "its floor to sit below " + str(ALPHA) + "."
        )
    lines.append("")

    lines.append("## Omnibus results (T82.5)")
    lines.append("")
    if omnibus is not None and len(omnibus):
        for _, row in omnibus[omnibus["metric"] == metric].iterrows():
            lines.append(
                "- **"
                + str(row["run"])
                + "** -- Friedman over "
                + str(int(row["k_models"]))
                + " models on n="
                + str(int(row["n"]))
                + " folds: chi2="
                + format(float(row["statistic"]), ".3f")
                + ", p="
                + format(float(row["p_value"]), ".4g")
                + ", Kendall W="
                + format(float(row["effect_size"]), ".3f")
                + ". "
                + str(row["posthoc_note"])
                + " Average ranks (lower is better): "
                + str(row["average_ranks"])
                + "."
            )
    else:
        lines.append("- No omnibus test could be run.")
    lines.append("")

    if posthoc is not None and len(posthoc):
        part = posthoc[posthoc["metric"] == metric]
        separated = part[part["significant"].astype(bool)]
        lines.append("### Nemenyi post-hoc: pairs separated by more than the critical difference")
        lines.append("")
        if len(separated):
            for _, row in separated.iterrows():
                lines.append(
                    "- **"
                    + str(row["run"])
                    + "** -- "
                    + str(row["model_a"])
                    + " (rank "
                    + format(float(row["average_rank_a"]), ".2f")
                    + ") vs "
                    + str(row["model_b"])
                    + " (rank "
                    + format(float(row["average_rank_b"]), ".2f")
                    + "): gap "
                    + format(float(row["rank_difference"]), ".2f")
                    + " > CD "
                    + format(float(row["critical_difference"]), ".2f")
                    + "."
                )
        else:
            lines.append(
                "- None. The omnibus rejected, but no individual pair is "
                "separated by more than the critical difference -- which is an "
                "ordinary and often-misreported outcome."
            )
        lines.append("")

    lines.append("## Sources")
    lines.append("")
    for item in sources:
        lines.append("- `" + str(item) + "`")
    lines.append("")

    target = Path(path)
    ensure_dir(target.parent)
    target.write_text("\n".join(lines), encoding="utf-8", newline="\n")
    log.info("wrote %d line(s) -> %s", len(lines), target)
    return target
