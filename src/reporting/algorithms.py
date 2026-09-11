"""ALG-01 to ALG-20: the pseudocode deliverables, generated from the code they describe.

A pseudocode block in a thesis is a claim about an implementation, and it goes
stale the same way a hand-typed table does: someone changes the filter order in
``configs/signal.yaml`` and the algorithm on page 71 still says four. So nothing
in these algorithms is typed twice.

* **Every parameter is read at export time** -- from the configs and module
  constants the implementation itself runs with -- and substituted into the
  step templates. A template holding a bare number other than 0 or 1 fails
  ``tests/test_algorithms.py``; the numbers live in the parameter block, with the
  file and key they came from.
* **Every step names the function that performs it**, resolved to ``path:line``
  with :mod:`inspect` when the file is written. A reference to a function that
  has moved or been renamed is an export error, not a stale citation.
* **Where a step quotes the code** -- a formula, a rule, a guard -- the quoted
  fragment is asserted to occur in that function's source (:func:`_quoted`). A
  changed formula therefore breaks the export instead of silently diverging
  from the page.

What the templates themselves say -- the order of the steps and what each one
does -- is the one thing written by hand, and it is what the behavioural
spot-checks in the gate test against the running code (T98.6, T98.7).

Phase 98 declares ALG-01 to ALG-10; Phase 99 adds ALG-11 to ALG-20 to the same
catalogue.
"""

from __future__ import annotations

import csv
import hashlib
import importlib
import inspect
import re
import string
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from src.utils.logging_setup import get_logger

__all__ = [
    "ALGORITHMS",
    "Algorithm",
    "AlgorithmError",
    "Parameter",
    "Step",
    "algorithms_dir",
    "algorithm_for",
    "render_text",
    "resolve_reference",
    "write_algorithm",
    "write_algorithms",
]

log = get_logger("reporting.algorithms")

INDEX_FILENAME = "algorithm_index.csv"
INDEX_COLUMNS = [
    "alg_id",
    "title",
    "task",
    "objective",
    "n_steps",
    "n_parameters",
    "implements",
    "source_sha256",
    "txt",
    "docx",
    "generated_utc",
]
COMMAND = "python scripts/39_export_algorithms.py"


class AlgorithmError(RuntimeError):
    """An algorithm cannot be rendered truthfully against the current code."""


# ---------------------------------------------------------------------------
# the model
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Step:
    """One pseudocode line. ``depth`` is the nesting level; ``ref`` the implementer."""

    text: str
    depth: int = 0
    ref: str = ""


@dataclass(frozen=True)
class Parameter:
    """A value substituted into the steps, with where it was read from."""

    name: str
    value: str
    source: str


@dataclass(frozen=True)
class Algorithm:
    alg_id: str
    title: str
    task: str
    objective: str
    purpose: str
    inputs: tuple[str, ...]
    outputs: tuple[str, ...]
    steps: tuple[Step, ...]
    parameters: Callable[[], list[Parameter]]
    notes: tuple[str, ...] = field(default_factory=tuple)

    def slug(self) -> str:
        cleaned = re.sub(r"[^a-z0-9]+", "_", self.title.lower()).strip("_")
        return self.alg_id + "_" + cleaned

    def references(self) -> list[str]:
        """Every implementer named by a step, in first-use order."""
        seen: list[str] = []
        for step in self.steps:
            if step.ref and step.ref not in seen:
                seen.append(step.ref)
        return seen


# ---------------------------------------------------------------------------
# reading the code
# ---------------------------------------------------------------------------


def _root() -> Path:
    return Path(__file__).resolve().parents[2]


def _object(reference: str) -> Any:
    module_name, _, qualname = reference.partition(":")
    if not qualname:
        raise AlgorithmError("reference " + repr(reference) + " is not module:qualname")
    target: Any = importlib.import_module(module_name)
    for part in qualname.split("."):
        if not hasattr(target, part):
            raise AlgorithmError(reference + ": " + part + " does not exist")
        target = getattr(target, part)
    if isinstance(target, property):  # cite the getter, which is where the code is
        target = target.fget
    return inspect.unwrap(target)


def resolve_reference(reference: str) -> tuple[str, int]:
    """``module:qualname`` -> ``(repo-relative path, first line)``. Raises if gone."""
    target = _object(reference)
    try:
        path = Path(inspect.getsourcefile(target) or "")
        _, line = inspect.getsourcelines(target)
    except (TypeError, OSError) as error:
        raise AlgorithmError(reference + " has no source to cite: " + str(error)) from error
    try:
        relative = path.resolve().relative_to(_root()).as_posix()
    except ValueError as error:
        raise AlgorithmError(reference + " is not inside the repository") from error
    return relative, int(line)


def _quoted(reference: str, fragment: str) -> str:
    """Return ``fragment`` after asserting the implementer's source contains it.

    The one way a formula from the code can appear on the page: quoted, and
    checked against the function that computes it every time the file is
    written.
    """
    source = " ".join(inspect.getsource(_object(reference)).split())
    if " ".join(fragment.split()) not in source:
        raise AlgorithmError(
            reference + " no longer contains " + repr(fragment) + "; the algorithm "
            "quoting it would describe code that does not exist"
        )
    return fragment


def _cfg(name: str, key: str) -> Any:
    from src.utils.config import load_config

    return load_config(name).require(key)


def _p(name: str, value: Any, source: str) -> Parameter:
    return Parameter(name, _show(value), source)


def _show(value: Any) -> str:
    if isinstance(value, bool):
        return str(value).lower()
    if value is None:
        return "none"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (list, tuple)):
        return ", ".join(_show(item) for item in value)
    return str(value)


def _space(model_id: str, name: str) -> str:
    """A search dimension from configs/models.yaml, written as the space it is."""
    spec = dict(_cfg("models", "models." + model_id + ".search_space." + name))
    kind = str(spec["type"])
    if kind == "categorical":
        return "{" + _show(spec["choices"]) + "}"
    scale = {"log_uniform": "log-uniform", "int_uniform": "integer", "uniform": "uniform"}
    return scale.get(kind, kind) + " [" + _show(spec["low"]) + ", " + _show(spec["high"]) + "]"


# ---------------------------------------------------------------------------
# parameters, per algorithm
# ---------------------------------------------------------------------------


def _params_alg01() -> list[Parameter]:
    from src.data_loader.catalog import DATASET_SHORT_NAMES, DATASET_TASKS

    return [
        _p("datasets", list(DATASET_SHORT_NAMES.values()), "src/data_loader/catalog.py"),
        _p(
            "tasks",
            [task for tasks in DATASET_TASKS.values() for task in tasks],
            "src/data_loader/catalog.py DATASET_TASKS",
        ),
        _p("uid_rule", "{dataset}_{subset}_{record_id}", "src/data_loader/master.py"),
    ]


def _params_alg02() -> list[Parameter]:
    native = dict(_cfg("signal", "resample.native_fs"))
    return [
        _p(
            "target_fs",
            _cfg("signal", "resample.target_fs"),
            "configs/signal.yaml resample.target_fs",
        ),
        _p("method", _cfg("signal", "resample.method"), "configs/signal.yaml resample.method"),
        _p("mono", _cfg("signal", "resample.mono"), "configs/signal.yaml resample.mono"),
        _p(
            "native_rates",
            [key + " " + _show(value) + " Hz" for key, value in native.items()],
            "configs/signal.yaml resample.native_fs",
        ),
        _p(
            "ratio_rule",
            _quoted("src.preprocessing.io:resample_ratio", "Fraction(int(fs_out), int(fs_in))"),
            "src/preprocessing/io.py resample_ratio",
        ),
    ]


def _params_alg03() -> list[Parameter]:
    from src.preprocessing.filters import default_padlen, design_bandpass

    target = int(_cfg("signal", "resample.target_fs"))
    low, high = _cfg("signal", "filter.low_hz"), _cfg("signal", "filter.high_hz")
    order = int(_cfg("signal", "filter.order"))
    zero_phase = bool(_cfg("signal", "filter.zero_phase"))
    padlen = default_padlen(design_bandpass(target, low, high, order))
    rule = _quoted(
        "src.preprocessing.filters:bandpass_filter",
        "effective_order = 2 * int(order) * (2 if zero_phase else 1)",
    )
    return [
        _p("target_fs", target, "configs/signal.yaml resample.target_fs"),
        _p("low_hz", low, "configs/signal.yaml filter.low_hz"),
        _p("high_hz", high, "configs/signal.yaml filter.high_hz"),
        _p("order", order, "configs/signal.yaml filter.order"),
        _p("zero_phase", zero_phase, "configs/signal.yaml filter.zero_phase"),
        _p("effective_order", 2 * order * (2 if zero_phase else 1), rule),
        _p("padlen", padlen, "filters.default_padlen(design_bandpass(...))"),
        _p(
            "min_padlen",
            _cfg("signal", "filter.min_padlen"),
            "configs/signal.yaml filter.min_padlen",
        ),
        _p(
            "norm_method",
            _cfg("signal", "normalization.method"),
            "configs/signal.yaml normalization.method",
        ),
        _p(
            "remove_dc",
            _cfg("signal", "normalization.remove_dc"),
            "configs/signal.yaml normalization.remove_dc",
        ),
        _p(
            "epsilon",
            _cfg("signal", "normalization.zero_variance_epsilon"),
            "configs/signal.yaml normalization.zero_variance_epsilon",
        ),
    ]


def _params_alg04() -> list[Parameter]:
    from src.feature_extraction.registry import (
        EXPECTED_FAMILY_COUNTS,
        EXPECTED_TOTAL,
        FAMILY_ORDER,
        registry_fingerprint,
    )

    return [
        _p("total", EXPECTED_TOTAL, "src/feature_extraction/registry.py EXPECTED_TOTAL"),
        _p(
            "families",
            [family + " " + str(EXPECTED_FAMILY_COUNTS[family]) for family in FAMILY_ORDER],
            "src/feature_extraction/registry.py FAMILY_ORDER, EXPECTED_FAMILY_COUNTS",
        ),
        _p("fingerprint", registry_fingerprint()[:16] + "...", "registry.registry_fingerprint()"),
        _p(
            "frame_length",
            _cfg("signal", "framing.frame_length"),
            "configs/signal.yaml framing.frame_length",
        ),
        _p(
            "hop_length",
            _cfg("signal", "framing.hop_length"),
            "configs/signal.yaml framing.hop_length",
        ),
        _p("window", _cfg("signal", "framing.window"), "configs/signal.yaml framing.window"),
    ]


def _params_alg05() -> list[Parameter]:
    scheme = dict(_cfg("experiments", "cv_schemes.repeated_5x5_grouped"))
    return [
        _p(
            "imputer",
            _cfg("models", "pipeline.imputer.strategy"),
            "configs/models.yaml pipeline.imputer.strategy",
        ),
        _p(
            "scaler",
            _cfg("models", "pipeline.scaler.kind"),
            "configs/models.yaml pipeline.scaler.kind",
        ),
        _p(
            "selector",
            _cfg("models", "pipeline.selector.enabled"),
            "configs/models.yaml pipeline.selector.enabled",
        ),
        _p(
            "class_weight",
            _cfg("models", "pipeline.class_weight.strategy"),
            "configs/models.yaml pipeline.class_weight.strategy",
        ),
        _p(
            "repeats",
            scheme["n_repeats"],
            "configs/experiments.yaml cv_schemes.repeated_5x5_grouped",
        ),
        _p(
            "splits", scheme["n_splits"], "configs/experiments.yaml cv_schemes.repeated_5x5_grouped"
        ),
        _p(
            "group_key",
            scheme["group_key"],
            "configs/experiments.yaml cv_schemes.repeated_5x5_grouped",
        ),
    ]


def _params_alg06() -> list[Parameter]:
    from src.optimization import multi_objective as mo

    settings = "configs/models.yaml optimization.feature_selection."
    weights = mo.load_weights()
    return [
        _p(
            "rankers",
            _cfg("models", "optimization.feature_selection.rankers"),
            settings + "rankers",
        ),
        _p("k_grid", _cfg("models", "optimization.feature_selection.k_grid"), settings + "k_grid"),
        _p(
            "model",
            _cfg("models", "optimization.feature_selection.evaluation_model"),
            settings + "evaluation_model",
        ),
        _p(
            "inner_splits",
            _cfg("models", "optimization.feature_selection.inner_splits"),
            settings + "inner_splits",
        ),
        _p(
            "repeats",
            _cfg("models", "optimization.feature_selection.repeats"),
            settings + "repeats",
        ),
        _p(
            "n_se",
            _cfg("models", "optimization.feature_selection.n_standard_errors"),
            settings + "n_standard_errors",
        ),
        _p("alpha", weights.alpha, "multi_objective.load_weights()"),
        _p("beta", weights.beta, "multi_objective.load_weights()"),
        _p("gamma", weights.gamma, "multi_objective.load_weights()"),
        _p("n_total", weights.n_features_total, "multi_objective.load_weights()"),
    ]


def _params_alg07() -> list[Parameter]:
    from src.evaluation.tuned import DEFAULT_MODEL_TRIALS, DEFAULT_TRIALS
    from src.optimization.base import DEFAULT_INNER_SPLITS
    from src.optimization.bayesian import ACQUISITION

    return [
        _p("trials", DEFAULT_TRIALS, "src/evaluation/tuned.py DEFAULT_TRIALS"),
        _p(
            "model_trials",
            [key + " " + str(value) for key, value in DEFAULT_MODEL_TRIALS.items()],
            "src/evaluation/tuned.py DEFAULT_MODEL_TRIALS",
        ),
        _p("inner_splits", DEFAULT_INNER_SPLITS, "src/optimization/base.py DEFAULT_INNER_SPLITS"),
        _p("acquisition", ACQUISITION, "src/optimization/bayesian.py ACQUISITION"),
        _p(
            "n_initial_rule",
            _quoted(
                "src.optimization.bayesian:BayesianSearch.__init__",
                "max(3, min(10, self.budget.max_trials // 4))",
            ),
            "src/optimization/bayesian.py BayesianSearch.__init__",
        ),
        _p(
            "binary_objective",
            _cfg("experiments", "defaults.scoring.binary"),
            "configs/experiments.yaml defaults.scoring",
        ),
        _p(
            "multiclass_objective",
            _cfg("experiments", "defaults.scoring.multiclass"),
            "configs/experiments.yaml defaults.scoring",
        ),
        _p("seed", _cfg("experiments", "defaults.seed"), "configs/experiments.yaml defaults.seed"),
    ]


def _params_alg08() -> list[Parameter]:
    from src.models.calibration import ISOTONIC_MIN_SAMPLES

    base = "configs/models.yaml models.M3."
    return [
        _p("kernel", _cfg("models", "models.M3.defaults.kernel"), base + "defaults.kernel"),
        _p("C", _cfg("models", "models.M3.defaults.C"), base + "defaults.C"),
        _p("gamma", _cfg("models", "models.M3.defaults.gamma"), base + "defaults.gamma"),
        _p(
            "class_weight",
            _cfg("models", "models.M3.defaults.class_weight"),
            base + "defaults.class_weight",
        ),
        _p("C_space", _space("M3", "C"), base + "search_space.C"),
        _p("gamma_space", _space("M3", "gamma"), base + "search_space.gamma"),
        _p("class_weight_space", _space("M3", "class_weight"), base + "search_space.class_weight"),
        _p(
            "method", _cfg("models", "calibration.method"), "configs/models.yaml calibration.method"
        ),
        _p("cv", _cfg("models", "calibration.cv"), "configs/models.yaml calibration.cv"),
        _p(
            "ensemble",
            _cfg("models", "calibration.ensemble"),
            "configs/models.yaml calibration.ensemble",
        ),
        _p("isotonic_min", ISOTONIC_MIN_SAMPLES, "src/models/calibration.py ISOTONIC_MIN_SAMPLES"),
    ]


def _calibrated_members() -> list[str]:
    return [str(m) for m in _cfg("models", "models.M6.calibrate_members")]


def _params_alg09() -> list[Parameter]:
    base = "configs/models.yaml models.M4."
    defaults = dict(_cfg("models", "models.M4.defaults"))
    return [
        _p("n_estimators", defaults["n_estimators"], base + "defaults.n_estimators"),
        _p("max_depth", defaults["max_depth"], base + "defaults.max_depth"),
        _p("min_samples_leaf", defaults["min_samples_leaf"], base + "defaults.min_samples_leaf"),
        _p("max_features", defaults["max_features"], base + "defaults.max_features"),
        _p("class_weight", defaults["class_weight"], base + "defaults.class_weight"),
        _p("seed", defaults["random_state"], base + "defaults.random_state"),
        _p("n_estimators_space", _space("M4", "n_estimators"), base + "search_space"),
        _p("max_depth_space", _space("M4", "max_depth"), base + "search_space"),
        _p("max_features_space", _space("M4", "max_features"), base + "search_space"),
        _p(
            "calibrated",
            "M4" in _calibrated_members(),
            "whether M4 is listed in configs/models.yaml models.M6.calibrate_members",
        ),
    ]


def _params_alg10() -> list[Parameter]:
    base = "configs/models.yaml models.M5."
    defaults = dict(_cfg("models", "models.M5.defaults"))
    return [
        _p("implementation", _cfg("models", "models.M5.implementation"), base + "implementation"),
        _p("n_estimators", defaults["n_estimators"], base + "defaults.n_estimators"),
        _p("learning_rate", defaults["learning_rate"], base + "defaults.learning_rate"),
        _p("max_depth", defaults["max_depth"], base + "defaults.max_depth"),
        _p("subsample", defaults["subsample"], base + "defaults.subsample"),
        _p("min_samples_leaf", defaults["min_samples_leaf"], base + "defaults.min_samples_leaf"),
        _p("class_weight", defaults["class_weight"], base + "defaults.class_weight"),
        _p(
            "weight_rule",
            _quoted(
                "src.models.weighting:balanced_sample_weight",
                "targets.size / (classes.size * counts.astype(float))",
            ),
            "src/models/weighting.py balanced_sample_weight",
        ),
        _p("learning_rate_space", _space("M5", "learning_rate"), base + "search_space"),
        _p("max_depth_space", _space("M5", "max_depth"), base + "search_space"),
        _p(
            "calibrated",
            "M5" in _calibrated_members(),
            "whether M5 is listed in configs/models.yaml models.M6.calibrate_members",
        ),
    ]


# ---------------------------------------------------------------------------
# the catalogue
# ---------------------------------------------------------------------------

_DL = "src.data_loader."
_PP = "src.preprocessing."
_FX = "src.feature_extraction.extractor:"
_MP = "src.models.pipeline:"
_FS = "src.feature_selection."
_OB = "src.optimization.base:"

ALGORITHMS: tuple[Algorithm, ...] = (
    Algorithm(
        "ALG-01",
        "Dataset loading and master metadata construction",
        "T98.1",
        "O1, O6",
        "Four public corpora become one table with one row per recording, each label "
        "written only into the task its corpus owns.",
        ("the four corpus folders named in configs/paths.yaml",),
        ("metadata_master.csv and .parquet (DA-08), one row per recording",),
        (
            Step(
                "D1 <- read every PhysioNet header and REFERENCE label; map each validation "
                "copy to its training twin",
                ref=_DL + "physionet:load_physionet",
            ),
            Step(
                "D2, D3 <- read PASCAL set_a and set_b; take labels from the Heartbeat_Sound "
                "folder structure, derive subject ids from filenames or set "
                "subject_derived <- false",
                ref=_DL + "pascal:load_pascal",
            ),
            Step(
                "D4 <- read CirCor; parse #Outcome from each patient .txt file; patient_id is "
                "native",
                ref=_DL + "circor:load_circor",
            ),
            Step(
                "catalog <- concatenate D1..D4 ({datasets}); each corpus writes only the class "
                "columns of the tasks it owns ({tasks})",
                ref=_DL + "catalog:build_catalog",
            ),
            Step(
                "assert no record carries a label of a task its corpus does not own",
                depth=1,
                ref=_DL + "catalog:build_catalog",
            ),
            Step(
                "scan <- decode every file; record silence, clipping and a content hash",
                ref=_DL + "integrity:scan_corpus",
            ),
            Step(
                "for each group of exact, content or near duplicates:",
                ref=_DL + "duplicates:build_duplicate_report",
            ),
            Step(
                "keep one member: never a Heartbeat_Sound entry, never a validation copy "
                "over its training twin, otherwise the smallest record_uid; mark the rest "
                "is_duplicate",
                depth=1,
                ref=_DL + "duplicates:keep_priority",
            ),
            Step(
                "master <- catalog + task labels + quality flags + duplicate decision",
                ref=_DL + "master:build_master",
            ),
            Step(
                "split_group <- subject_id if use_in_supervised else blank",
                depth=1,
                ref=_DL + "master:build_master",
            ),
            Step(
                "validate: schema order and dtypes, record_uid = {uid_rule} and unique, "
                "duration > 0, split_group present exactly for supervised rows, no "
                "cross-task label",
                ref=_DL + "master:validate_master",
            ),
            Step("write master as CSV and Parquet", ref=_DL + "master:write_master"),
        ),
        _params_alg01,
        notes=(
            "Heartbeat_Sound is a byte-identical copy of set_a and set_b and is used only as "
            "the label source; including it as data would double those corpora.",
        ),
    ),
    Algorithm(
        "ALG-02",
        "Signal harmonization and resampling",
        "T98.1",
        "O4",
        "Every recording, whatever its native rate and channel count, becomes a mono "
        "float signal at one common rate, and its quality is measured before filtering.",
        ("one WAV file",),
        ("mono signal at {target_fs} Hz, its native rate, quality measurements",),
        (
            Step(
                "(x, fs_native) <- decode WAV as float32 scaled to [-1, 1)", ref=_PP + "io:load_wav"
            ),
            Step(
                "if x has more than one channel: x <- channel {mono}; else unchanged",
                ref=_PP + "io:to_mono",
            ),
            Step("if fs_native = {target_fs} or x is empty: y <- x", ref=_PP + "io:resample_to"),
            Step(
                "else: y <- band-limited resample of x from fs_native to {target_fs} Hz "
                "with {method}; exact rational ratio {ratio_rule}",
                depth=0,
                ref=_PP + "io:resample_to",
            ),
            Step(
                "quality <- clipping, silence, SNR proxy and drift ratio measured on y, "
                "before the band-pass removes the out-of-band term they need",
                ref=_PP + "quality:measure_signal",
            ),
            Step(
                "return y, fs_native, quality and the list of steps applied",
                ref=_PP + "pipeline:preprocess",
            ),
        ),
        _params_alg02,
        notes=("Native rates in this corpus: {native_rates}.",),
    ),
    Algorithm(
        "ALG-03",
        "Butterworth filtering and z-score normalization",
        "T98.2",
        "O4",
        "A zero-phase band-pass keeps the heart-sound band; a z-score removes loudness "
        "differences between stethoscopes, gains and chest walls.",
        ("signal y at {target_fs} Hz from ALG-02",),
        ("filtered, normalized float32 signal z, with what was applied",),
        (
            Step(
                "sos <- Butterworth band-pass, prototype order {order}, {low_hz}-{high_hz} Hz "
                "at {target_fs} Hz, as second-order sections (designed once, cached)",
                ref=_PP + "filters:design_bandpass",
            ),
            Step("if the filter is disabled: x_f <- y", ref=_PP + "filters:bandpass_filter"),
            Step("else:", ref=_PP + "filters:bandpass_filter"),
            Step(
                "padlen <- {padlen} samples, sosfiltfilt's own default for this design",
                depth=1,
                ref=_PP + "filters:default_padlen",
            ),
            Step(
                "if |y| <= padlen: padlen <- |y| - 1", depth=1, ref=_PP + "filters:bandpass_filter"
            ),
            Step(
                "if padlen < {min_padlen}: x_f <- y, flag too_short",
                depth=2,
                ref=_PP + "filters:bandpass_filter",
            ),
            Step(
                "x_f <- sosfiltfilt(sos, y, padlen): forward then backward, zero phase, "
                "effective order {effective_order}",
                depth=1,
                ref=_PP + "filters:bandpass_filter",
            ),
            Step(
                "if normalization is disabled or method is none: return x_f",
                ref=_PP + "normalize:normalize",
            ),
            Step(
                "if remove_dc ({remove_dc}): x_f <- x_f - mean(x_f)",
                ref=_PP + "normalize:remove_dc",
            ),
            Step("sigma <- std(x_f) in float64", ref=_PP + "normalize:zscore_normalize"),
            Step(
                "if sigma <= {epsilon}: return the input unchanged, flag zero_variance",
                ref=_PP + "normalize:zscore_normalize",
            ),
            Step(
                "z <- (x_f - mean(x_f)) / sigma as float32 (method {norm_method})",
                ref=_PP + "normalize:zscore_normalize",
            ),
        ),
        _params_alg03,
        notes=(
            "The quoted order is the prototype's: the band-pass transform doubles it and the "
            "forward-backward pass doubles it again, to {effective_order}. The two passes "
            "square the magnitude response, so each cutoff is attenuated twice as many dB as "
            "a single pass would attenuate it.",
        ),
    ),
    Algorithm(
        "ALG-04",
        "Multi-domain 138-feature extraction",
        "T98.2",
        "O4",
        "One preprocessed signal becomes one fixed-order vector of {total} values from "
        "six feature families, the same columns for every corpus.",
        ("signal z from ALG-03 and its rate",),
        ("vector of {total} values in registry order, flags, failed families, timings",),
        (
            Step("limit BLAS to one thread for the whole extraction", ref=_FX + "extract_all"),
            Step(
                "for each family f in registry order ({families}):", ref=_FX + "_extract_all_inner"
            ),
            Step(
                "if f was not requested: its columns <- NaN, flag f_not_run",
                depth=1,
                ref=_FX + "_extract_all_inner",
            ),
            Step(
                "r <- extractor_f(z, fs); frame-based families use {frame_length}-sample "
                "frames, hop {hop_length}, {window} window",
                depth=1,
                ref="src.feature_extraction.registry:get_extractor",
            ),
            Step(
                "assert |r| = the family's registered count",
                depth=1,
                ref=_FX + "_extract_all_inner",
            ),
            Step(
                "append r; keep its flags and seconds; a failed family keeps NaN in its "
                "own columns and its error message",
                depth=1,
                ref=_FX + "_extract_all_inner",
            ),
            Step(
                "assert the vector has {total} values whose names equal the registry order "
                "exactly (fingerprint {fingerprint})",
                ref=_FX + "_validate",
            ),
            Step(
                "return the vector, flags, failed families and per-family seconds",
                ref=_FX + "extract_all",
            ),
        ),
        _params_alg04,
        notes=(
            "Thread count is pinned because a threaded BLAS changes the last bit of the "
            "chroma features between runs, and the fixed-seed rule requires identical numbers.",
        ),
    ),
    Algorithm(
        "ALG-05",
        "Fold-safe scaling and preprocessing",
        "T98.3",
        "O1",
        "Everything that learns from data -- imputation, scaling, selection, class "
        "weighting -- is fitted on the training fold only, as steps of one pipeline.",
        ("feature matrix X, labels y, the stored fold map (DA-07)",),
        ("per-fold out-of-fold predictions from a pipeline that never saw its test rows",),
        (
            Step(
                "for each (repeat, fold) of the stored map ({repeats} x {splits}, grouped on "
                "{group_key}), loaded never re-derived:",
                ref=_DL + "splits:iter_folds",
            ),
            Step(
                "assert no subject appears in both train and test",
                depth=1,
                ref=_DL + "splits:assert_no_leakage",
            ),
            Step(
                "P <- Pipeline[imputer ({imputer}), scaler ({scaler}), selector if enabled "
                "({selector}), clone(estimator)]",
                depth=1,
                ref=_MP + "build_pipeline",
            ),
            Step(
                "if the estimator accepts class_weight: set it to {class_weight}; otherwise "
                "leave imbalance to the estimator (ALG-10's wrapper)",
                depth=1,
                ref=_MP + "class_weight_for",
            ),
            Step(
                "if resampling is enabled: refuse -- it must happen inside the fold",
                depth=1,
                ref=_MP + "build_pipeline",
            ),
            Step(
                "fit P on X[train], y[train]: the medians, means, variances and selector "
                "scores are learned from training rows only",
                depth=1,
                ref=_MP + "build_pipeline",
            ),
            Step(
                "predict X[test] through P: the test rows are transformed, never fitted",
                depth=1,
                ref=_MP + "build_pipeline",
            ),
            Step(
                "expose the learned statistics for the fold-safety tests", ref=_MP + "fitted_steps"
            ),
        ),
        _params_alg05,
    ),
    Algorithm(
        "ALG-06",
        "Search based feature selection",
        "T98.3",
        "O5",
        "SO-04: rank features inside inner training folds, cut every ranking at every k, "
        "and choose the cheapest configuration statistically indistinguishable from the best.",
        ("X, y, subject groups, outer folds of repeat(s) {repeats}",),
        (
            "the chosen (ranker, k), per-fold subsets, the consensus subset FE-12, and the "
            "all-features-versus-subset comparison",
        ),
        (
            Step("for each outer fold o:", ref=_FS + "sweep:sweep_feature_counts"),
            Step(
                "inner <- {inner_splits} grouped, stratified splits of o's training rows",
                depth=1,
                ref=_OB + "inner_folds",
            ),
            Step(
                "for each inner split (T, V) and each ranker r in {rankers}:",
                depth=1,
                ref=_FS + "sweep:sweep_feature_counts",
            ),
            Step(
                "rank <- r fitted on median-imputed X[T], y[T]; ties broken by column position",
                depth=2,
                ref=_FS + "ranking:rank_features",
            ),
            Step(
                "for each k in {k_grid}: S <- top k; fit {model} on X[T, S] in the fold-safe "
                "pipeline; score V",
                depth=2,
                ref=_FS + "sweep:score_subset",
            ),
            Step(
                "J <- alpha (1 - macroF1) + beta k / N + gamma t(S), alpha = {alpha}, "
                "beta = {beta}, gamma = {gamma}, N = {n_total}; t(S) is the extraction time "
                "of the families S still needs, normalised by all features",
                depth=3,
                ref="src.optimization.multi_objective:score_j",
            ),
            Step(
                "for each (r, k): mean J, mean macro-F1 and its standard error over all "
                "inner evaluations",
                ref=_FS + "sweep:choose_configuration",
            ),
            Step(
                "guard <- best mean macro-F1 - {n_se} x its standard error",
                ref=_FS + "sweep:choose_configuration",
            ),
            Step(
                "(r*, k*) <- lowest mean J among configurations with mean macro-F1 >= guard; "
                "ties to the smaller k, then the ranker name",
                ref=_FS + "sweep:choose_configuration",
            ),
            Step(
                "for each outer fold: re-rank with r* on that fold's training rows; keep the "
                "top k*",
                ref=_FS + "sweep:select_per_fold",
            ),
            Step(
                "FE-12 <- the k* features chosen in the most folds; ties by column position",
                ref=_FS + "sweep:consensus_subset",
            ),
            Step(
                "score all features and the per-fold subsets on the same outer test folds",
                ref=_FS + "sweep:compare_all_versus_selected",
            ),
        ),
        _params_alg06,
        notes=(
            "The guard is a selection rule fixed before the result; it is never tuned to a score.",
        ),
    ),
    Algorithm(
        "ALG-07",
        "Hyperparameter optimization",
        "T98.4",
        "O5",
        "Nested Bayesian search: each outer fold's hyperparameters are chosen on inner "
        "splits of its own training rows, and the outer test rows are never scored.",
        ("X, y, groups, one model's declared search space, the outer fold map",),
        ("per outer fold: the chosen point, the trial log, and one outer score",),
        (
            Step("for each outer fold o:", ref="src.evaluation.tuned:NestedSearchPlanner.plan"),
            Step(
                "splits <- {inner_splits} grouped, stratified inner splits of o's training rows",
                depth=1,
                ref=_OB + "inner_folds",
            ),
            Step(
                "budget <- {trials} trials per model (exceptions: {model_trials}); a "
                "wall-clock ceiling, if set, is checked between trials",
                depth=1,
                ref="src.evaluation.tuned:NestedSearchPlanner.budget_for",
            ),
            Step(
                "surrogate <- Gaussian process, acquisition {acquisition}, initial random "
                "points {n_initial_rule}, seed {seed}",
                depth=1,
                ref="src.optimization.bayesian:BayesianSearch.__init__",
            ),
            Step("while the budget allows another trial:", depth=1, ref=_OB + "BaseSearch.run"),
            Step(
                "theta <- decode(ask()); repair it into the legal space; a violated "
                "constraint is logged as an invalid trial",
                depth=2,
                ref=_OB + "BaseSearch.run",
            ),
            Step(
                "for each inner split (T, V): fit a fresh pipeline of model(theta) on T, "
                "predict V, record T and V in the row ledger",
                depth=2,
                ref=_OB + "score_params",
            ),
            Step(
                "score <- mean over inner splits of the objective ({binary_objective} for "
                "binary, {multiclass_objective} for multiclass); a fit that raises is a "
                "failed trial",
                depth=2,
                ref=_OB + "score_params",
            ),
            Step(
                "tell(theta, -score); a failed trial is told the worst attainable value",
                depth=2,
                ref="src.optimization.bayesian:BayesianSearch._observe",
            ),
            Step(
                "theta* <- the best trial; ties to the earlier trial",
                depth=1,
                ref=_OB + "SearchResult.best",
            ),
            Step(
                "assert the ledger holds no row of o's test fold",
                depth=1,
                ref=_OB + "RowLedger.assert_disjoint_from",
            ),
            Step(
                "fit model(theta*) on all of o's training rows; score o's test rows once",
                depth=1,
                ref="src.evaluation.tuned:NestedSearchPlanner.plan",
            ),
        ),
        _params_alg07,
        notes=(
            "For the ensembles M6 and M7 the members are what is searched: they declare no "
            "hyperparameters of their own.",
        ),
    ),
    Algorithm(
        "ALG-08",
        "SVM training and probability calibration",
        "T98.4",
        "O3",
        "M3: an RBF support vector machine whose scores are turned into probabilities by an "
        "explicit, visible calibration fitted inside the training fold.",
        ("training rows X, y (optionally their subject groups)",),
        ("calibrated class probabilities whose argmax is the prediction",),
        (
            Step(
                "svc <- SVC(kernel {kernel}, C {C}, gamma {gamma}, class_weight "
                "{class_weight}); probability is never enabled",
                ref="src.models.estimators:make_m3",
            ),
            Step(
                "M3 <- CalibratedSVM(svc, method {method}, cv {cv}, ensemble {ensemble})",
                ref="src.models.estimators:make_m3",
            ),
            Step("fit(X, y):", ref="src.models.calibration:CalibratedSVM.fit"),
            Step(
                "refuse an inner estimator with probability enabled; forward class_weight "
                "to the inner SVC",
                depth=1,
                ref="src.models.calibration:CalibratedSVM.fit",
            ),
            Step(
                "if method is isotonic and |y| < {isotonic_min}: warn that it overfits",
                depth=1,
                ref="src.models.calibration:CalibratedSVM.fit",
            ),
            Step(
                "split the training rows into {cv} calibration folds, subject-grouped when a "
                "grouped split is supplied",
                depth=1,
                ref="src.models.calibration:grouped_calibration_cv",
            ),
            Step(
                "for each calibration fold: fit svc on its training part; fit the sigmoid "
                "p = 1 / (1 + exp(A f(x) + B)) to svc's decision values on its held part",
                depth=1,
                ref="src.models.calibration:CalibratedSVM.fit",
            ),
            Step(
                "keep every (svc, sigmoid) pair; predict_proba <- their mean probability",
                depth=1,
                ref="src.models.calibration:CalibratedSVM.predict_proba",
            ),
            Step(
                "predict <- argmax of predict_proba, so the two can never disagree",
                ref="src.models.calibration:CalibratedSVM.predict",
            ),
        ),
        _params_alg08,
        notes=(
            "Searched space: C {C_space}, gamma {gamma_space}, class_weight {class_weight_space}.",
        ),
    ),
    Algorithm(
        "ALG-09",
        "Random Forest training",
        "T98.5",
        "O3",
        "M4: a class-weighted random forest, the ensemble's bagged member and the source of "
        "one of the feature rankings.",
        ("training rows X, y",),
        ("class probabilities and impurity-based feature importances",),
        (
            Step(
                "RF <- RandomForestClassifier(n_estimators {n_estimators}, max_depth "
                "{max_depth}, min_samples_leaf {min_samples_leaf}, max_features "
                "{max_features}, random_state {seed})",
                ref="src.models.estimators:make_m4",
            ),
            Step(
                "class_weight <- {class_weight}, set inside the fold pipeline from the "
                "training labels",
                ref=_MP + "class_weight_for",
            ),
            Step(
                "for each tree t: draw a bootstrap sample of the training rows; grow the tree, "
                "choosing each split among {max_features} random candidate features by "
                "weighted Gini impurity, until min_samples_leaf or max_depth stops it",
                ref="src.models.estimators:make_m4",
            ),
            Step(
                "predict_proba <- mean over trees of each leaf's class-weighted class fractions",
                ref="src.models.estimators:make_m4",
            ),
            Step(
                "feature_importances_ <- mean impurity decrease per feature; read by the "
                "rf_importance ranker of ALG-06",
                ref=_FS + "ranking:_tree_importance",
            ),
            Step(
                "calibrated inside the ensemble: {calibrated}",
                ref="src.models.estimators:make_ensemble",
            ),
        ),
        _params_alg09,
        notes=(
            "Searched space: n_estimators {n_estimators_space}, max_depth {max_depth_space}, "
            "max_features {max_features_space}. Each tree is seeded from random_state, so the "
            "thread count does not change the forest.",
        ),
    ),
    Algorithm(
        "ALG-10",
        "Gradient Boosting training",
        "T98.5",
        "O3",
        "M5: gradient-boosted trees, given class weights through per-row sample weights "
        "because the estimator has no class_weight of its own.",
        ("training rows X, y",),
        ("class probabilities from an additive model of shallow trees",),
        (
            Step(
                "GB <- GradientBoostingClassifier(n_estimators {n_estimators}, learning_rate "
                "{learning_rate}, max_depth {max_depth}, subsample {subsample}, "
                "min_samples_leaf {min_samples_leaf}) ({implementation} implementation)",
                ref="src.models.estimators:make_m5",
            ),
            Step(
                "M5 <- ClassWeightedClassifier(GB, class_weight {class_weight})",
                ref="src.models.estimators:make_m5",
            ),
            Step("fit(X, y):", ref="src.models.weighting:ClassWeightedClassifier.fit"),
            Step(
                "w_i <- {weight_rule}, evaluated at y_i's class (balanced weighting)",
                depth=1,
                ref="src.models.weighting:balanced_sample_weight",
            ),
            Step(
                "if sample weights are supplied: w <- w x supplied",
                depth=1,
                ref="src.models.weighting:ClassWeightedClassifier.fit",
            ),
            Step(
                "F_0 <- log-odds of the weighted class prior",
                depth=1,
                ref="src.models.weighting:ClassWeightedClassifier.fit",
            ),
            Step(
                "for each boosting stage m: fit a regression tree of depth <= {max_depth} to "
                "the weighted negative gradient on a {subsample} fraction of rows; "
                "F_m <- F_(m-1) + {learning_rate} x tree",
                depth=1,
                ref="src.models.weighting:ClassWeightedClassifier.fit",
            ),
            Step(
                "predict_proba <- sigmoid(F) for two classes, softmax over per-class F otherwise",
                ref="src.models.weighting:ClassWeightedClassifier.predict_proba",
            ),
            Step(
                "calibrated inside the ensemble: {calibrated}",
                ref="src.models.estimators:make_ensemble",
            ),
        ),
        _params_alg10,
        notes=(
            "Searched space: learning_rate {learning_rate_space}, max_depth "
            "{max_depth_space}. Without the wrapper M5 would be the only ensemble member "
            "with no imbalance handling at all.",
        ),
    ),
)

_BY_ID = {algorithm.alg_id: algorithm for algorithm in ALGORITHMS}


def algorithm_for(alg_id: str) -> Algorithm:
    if alg_id not in _BY_ID:
        raise KeyError(alg_id + " is not in the algorithm catalogue")
    return _BY_ID[alg_id]


# ---------------------------------------------------------------------------
# rendering
# ---------------------------------------------------------------------------


def _fill(template: str, values: dict[str, str], where: str) -> str:
    names = {name for _, name, _, _ in string.Formatter().parse(template) if name}
    missing = sorted(names - set(values))
    if missing:
        raise AlgorithmError(where + " uses unknown parameter(s): " + ", ".join(missing))
    return template.format_map(values)


@dataclass(frozen=True)
class Rendered:
    """An algorithm with its parameters substituted and its references resolved."""

    algorithm: Algorithm
    parameters: tuple[Parameter, ...]
    purpose: str
    inputs: tuple[str, ...]
    outputs: tuple[str, ...]
    lines: tuple[tuple[int, str, str], ...]  # (depth, text, "path:line")
    notes: tuple[str, ...]
    implements: tuple[tuple[str, str, int], ...]  # (reference, path, line)


def render(algorithm: Algorithm) -> Rendered:
    parameters = tuple(algorithm.parameters())
    values = {parameter.name: parameter.value for parameter in parameters}
    where = algorithm.alg_id
    implements: list[tuple[str, str, int]] = []
    located: dict[str, str] = {}
    for reference in algorithm.references():
        path, line = resolve_reference(reference)
        implements.append((reference, path, line))
        located[reference] = path + ":" + str(line)
    lines = tuple(
        (step.depth, _fill(step.text, values, where), located.get(step.ref, ""))
        for step in algorithm.steps
    )
    return Rendered(
        algorithm=algorithm,
        parameters=parameters,
        purpose=_fill(algorithm.purpose, values, where),
        inputs=tuple(_fill(text, values, where) for text in algorithm.inputs),
        outputs=tuple(_fill(text, values, where) for text in algorithm.outputs),
        lines=lines,
        notes=tuple(_fill(text, values, where) for text in algorithm.notes),
        implements=tuple(implements),
    )


def _step_lines(rendered: Rendered) -> list[str]:
    width = len(str(len(rendered.lines)))
    return [
        str(number).rjust(width) + "  " + "    " * depth + text
        for number, (depth, text, _) in enumerate(rendered.lines, start=1)
    ]


def render_text(rendered: Rendered) -> str:
    algorithm = rendered.algorithm
    name_width = max((len(p.name) for p in rendered.parameters), default=0)
    value_width = max((len(p.value) for p in rendered.parameters), default=0)
    out = [
        "Algorithm " + algorithm.alg_id + ": " + algorithm.title,
        "PV-MEPCG / PulseVision  |  " + algorithm.task + "  |  objective " + algorithm.objective,
        "",
        "Purpose: " + rendered.purpose,
        "Input:   " + "; ".join(rendered.inputs),
        "Output:  " + "; ".join(rendered.outputs),
        "",
        "Parameters (read from the repository when this file was generated)",
    ]
    out += [
        "  " + p.name.ljust(name_width) + "  " + p.value.ljust(value_width) + "  " + p.source
        for p in rendered.parameters
    ]
    out += ["", "Algorithm"]
    out += ["  " + line for line in _step_lines(rendered)]
    if rendered.notes:
        out += ["", "Notes"]
        out += ["  " + note for note in rendered.notes]
    out += ["", "Implemented by"]
    out += ["  " + path + ":" + str(line) + "  " + ref for ref, path, line in rendered.implements]
    out += ["", "Regenerate: " + COMMAND + " --ids " + algorithm.alg_id, ""]
    return "\n".join(out)


def _write_docx(rendered: Rendered, target: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    algorithm = rendered.algorithm
    document = Document()

    heading = document.add_paragraph()
    run = heading.add_run("Algorithm " + algorithm.alg_id + ": " + algorithm.title)
    run.bold = True
    run.font.size = Pt(11)

    for label, text in (
        ("Purpose", rendered.purpose),
        ("Input", "; ".join(rendered.inputs)),
        ("Output", "; ".join(rendered.outputs)),
    ):
        paragraph = document.add_paragraph()
        key = paragraph.add_run(label + ": ")
        key.bold = True
        key.font.size = Pt(9)
        body = paragraph.add_run(text)
        body.font.size = Pt(9)

    grid = document.add_table(rows=1, cols=3)
    grid.style = "Table Grid"
    for cell, header in zip(grid.rows[0].cells, ("Parameter", "Value", "Read from"), strict=True):
        cell.text = header
        for cell_run in cell.paragraphs[0].runs:
            cell_run.bold = True
            cell_run.font.size = Pt(8)
    for parameter in rendered.parameters:
        cells = grid.add_row().cells
        for cell, value in zip(
            cells, (parameter.name, parameter.value, parameter.source), strict=True
        ):
            cell.text = value
            for cell_run in cell.paragraphs[0].runs:
                cell_run.font.size = Pt(8)

    document.add_paragraph()
    for line in _step_lines(rendered):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        code = paragraph.add_run(line)
        code.font.name = "Consolas"
        code.font.size = Pt(8.5)

    for note in rendered.notes:
        paragraph = document.add_paragraph(note)
        paragraph.runs[0].italic = True
        paragraph.runs[0].font.size = Pt(8)

    paragraph = document.add_paragraph("Implemented by:")
    paragraph.runs[0].font.size = Pt(8)
    for ref, path, lineno in rendered.implements:
        cited = document.add_paragraph(path + ":" + str(lineno) + "  " + ref)
        cited.paragraph_format.space_after = Pt(0)
        cited.runs[0].font.name = "Consolas"
        cited.runs[0].font.size = Pt(7)

    footer = document.add_paragraph(
        "PV-MEPCG / PulseVision  |  generated from the implementation  |  regenerate: "
        + COMMAND
        + " --ids "
        + algorithm.alg_id
    )
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(7)
    document.save(str(target))
    return target


# ---------------------------------------------------------------------------
# writing
# ---------------------------------------------------------------------------


def algorithms_dir(out_dir: str | Path | None = None) -> Path:
    if out_dir is not None:
        return Path(out_dir)
    return Path(_cfg("paths", "outputs.algorithms"))


def _source_digest(rendered: Rendered) -> str:
    digest = hashlib.sha256()
    for path in sorted({path for _, path, _ in rendered.implements}):
        digest.update(path.encode("utf-8"))
        digest.update((_root() / path).read_bytes())
    return digest.hexdigest()


def _update_index(row: dict[str, str], target: Path) -> None:
    rows: list[dict[str, str]] = []
    if target.is_file():
        with target.open("r", encoding="utf-8", newline="") as handle:
            rows = [dict(item) for item in csv.DictReader(handle)]
    rows = [item for item in rows if item.get("alg_id") != row["alg_id"]] + [row]
    rows.sort(key=lambda item: item["alg_id"])
    with target.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=INDEX_COLUMNS)
        writer.writeheader()
        for item in rows:
            writer.writerow({column: item.get(column, "") for column in INDEX_COLUMNS})


def write_algorithm(
    algorithm: Algorithm,
    out_dir: str | Path | None = None,
    *,
    evidence_index: str | Path | None = None,
) -> dict[str, Path]:
    """Render one algorithm to plain text and DOCX, index it, register the evidence."""
    from src.utils.evidence import register_evidence
    from src.utils.io import ensure_dir

    rendered = render(algorithm)
    target_dir = Path(ensure_dir(algorithms_dir(out_dir)))
    txt = target_dir / (algorithm.slug() + ".txt")
    txt.write_text(render_text(rendered), encoding="utf-8")
    docx = _write_docx(rendered, target_dir / (algorithm.slug() + ".docx"))

    _update_index(
        {
            "alg_id": algorithm.alg_id,
            "title": algorithm.title,
            "task": algorithm.task,
            "objective": algorithm.objective,
            "n_steps": str(len(rendered.lines)),
            "n_parameters": str(len(rendered.parameters)),
            "implements": "; ".join(
                path + ":" + str(line) for _, path, line in rendered.implements
            ),
            "source_sha256": _source_digest(rendered),
            "txt": txt.name,
            "docx": docx.name,
            "generated_utc": datetime.now(UTC).isoformat(),
        },
        target_dir / INDEX_FILENAME,
    )
    register_evidence(
        algorithm.alg_id,
        txt,
        metric_or_asset=algorithm.title,
        objective=algorithm.objective,
        source_data="; ".join(sorted({path for _, path, _ in rendered.implements})),
        command=COMMAND + " --ids " + algorithm.alg_id,
        index_path=evidence_index,
    )
    log.info("%s: %s, %s", algorithm.alg_id, txt.name, docx.name)
    return {"txt": txt, "docx": docx}


def write_algorithms(
    ids: list[str] | None = None,
    out_dir: str | Path | None = None,
    *,
    evidence_index: str | Path | None = None,
) -> dict[str, dict[str, Path]]:
    wanted = ids or [algorithm.alg_id for algorithm in ALGORITHMS]
    return {
        alg_id: write_algorithm(algorithm_for(alg_id), out_dir, evidence_index=evidence_index)
        for alg_id in wanted
    }
